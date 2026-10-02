"""Validates: REQ-DAILY-001, REQ-DAILY-006, REQ-DAILY-013, REQ-DAILY-018, REQ-DAILY-022 (TEST-DAILY-001),
REQ-QAINTEL-003, REQ-QAINTEL-018, REQ-QAINTEL-023.

개편 전 SRS 스냅샷 위치(`data/daily_qa/snapshots/<날짜>.json`)를 옛 제품의 비교 기준으로 읽는 것까지 본다.
실제 Polarion·Claude·메일을 부르지 않는다.
"""

from __future__ import annotations

import json
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.modules.daily_qa.agent_runner import FakeRunner
from app.modules.daily_qa.pipeline import Inputs, RunLocked, run_daily
from app.modules.daily_qa.rules import RulesState
from app.modules.daily_qa.schema import SKILL_COVERAGE, SKILL_E, SKILL_F, SKILL_SPEC_SUMMARY
from app.modules.daily_qa.settings import PolarionSettings
from app.modules.daily_qa.srs_snapshot import save_snapshot
from app.modules.daily_qa.store import DailyQaStore
from tests.daily_qa_fixtures import FakePolarion, issue_item, make_settings, make_tc_workbook, rules_ok, srs_item, write_result

MONDAY = datetime(2026, 9, 28, 9, 0, tzinfo=timezone.utc)
TUESDAY = datetime(2026, 9, 29, 9, 0, tzinfo=timezone.utc)
YESTERDAY_SRS = [
    {"id": "VP-10", "old_id": "01-10-10", "title": "Viewer", "status": "draft", "text": "옛 본문", "is_category": False},
    {"id": "VP-11", "old_id": "01-10-11", "title": "Setting", "status": "draft", "text": "본문", "is_category": False},
]


@pytest.fixture
def env(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    cfg = make_settings(root, tmp_path / "ws")
    tc = make_tc_workbook(tmp_path / "(TC) 영향성평가_Checklist.xlsx", [("TC_1", "VP-10", "A"), ("TC_2", "VP-11", "B")])
    inputs = Inputs(tc_paths=[tc], manuals={"Operation Manual.txt": "1. 설치\n\n2. Viewer 화면에서 목록을 본다."}, checklist_template=tc)
    store = DailyQaStore(cfg.db_path)
    # 개편 전 위치의 어제 스냅샷: VP-10 본문이 오늘과 다르다 → Coverage 분석이 생긴다
    save_snapshot(root / "data" / "daily_qa" / "snapshots", "2026-09-27", YESTERDAY_SRS)
    polarion = FakePolarion(
        srs=[srs_item("VP-10", "01-10-10", "Viewer", text="새 본문"), srs_item("VP-11", "01-10-11", "Setting"),
             srs_item("VP-12", "01-10-12", "Report")],
        issues=[issue_item("VP-6669", "2026-09-28T01:00:00Z", ["VP-10"])],
    )
    sent = []
    return {
        "cfg": cfg, "inputs": inputs, "store": store, "polarion": polarion, "sent": sent, "tmp": tmp_path,
        "sender": lambda subject, text, body, to: sent.append((subject, text, to)) or {"status": "sent"},
    }


def _producer(task, run):
    payload = json.loads((run.in_dir / f"{task.task_id}.json").read_text(encoding="utf-8"))
    findings = []
    for item in payload.get("items", []):
        target = item["target"]
        findings.append({
            "subject": target, "verdict": "NOT_COVERED", "summary": "바뀐 사양을 검증하는 TC 없음",
            "evidence": [{"source_type": "srs", "location": f"{target} / 본문", "validity": "Current"}],
            "draft_tcs": [{"kind": "Regression", "srs_no": target, "title": f"{target} 새 동작", "test_step": "1. 실행한다.",
                           "expected_result": "1. 새 본문대로 표시된다.", "perspective": "DIRECT_SPEC", "rationale": "기존 TC 없음"}],
            "sections": {"change": {"summary": "본문 변경"},
                         "checklist_coverage": {"tcs": [{"tc_id": "TC_1", "location": "", "decision": "UPDATE_EXISTING",
                                                         "reason": "Expected 가 옛 본문", "recommended_change": "Expected 1번 수정"}]}},
        })
    if task.skill == SKILL_SPEC_SUMMARY:
        # 자동 사양 변경 분석은 TC 를 보지 않는다. 판정·요약만 쓴다 (REQ-QAINTEL-031).
        findings = [{**finding, "verdict": "QA_CHECK_NEEDED", "draft_tcs": [], "sections": {"change": {"summary": "본문 변경"}}}
                    for finding in findings]
    write_result(run, task, findings, questions=[{"subject": "VP-10", "question": "Admin 만 보이나?", "why": "권한"}]
                 if task.skill == SKILL_COVERAGE else None)


def _run(env, **kwargs):
    kwargs.setdefault("today", MONDAY)
    kwargs.setdefault("runner", FakeRunner(_producer))
    kwargs.setdefault("rules_state", rules_ok(env["tmp"]))
    return run_daily(env["cfg"], polarion_factory=lambda: env["polarion"], inputs=env["inputs"], send_email=env["sender"],
                     store=env["store"], spec_chunks_loader=lambda: ([], {}, []), **kwargs)


def test_full_run_uses_legacy_snapshot_and_creates_findings_draft_and_email(env):
    outcome = _run(env)
    stages = outcome["stages"]
    assert outcome["status"] == "SUCCESS", stages
    assert stages["collect_srs"]["counts"] == {"total": 3, "added": 1, "removed": 0, "modified": 1}
    assert stages["collect_issues"]["status"] == "ok" and "기준 스냅샷만" in stages["collect_issues"]["note"]
    assert stages["E"]["status"] == "ok" and stages["F"]["status"] == "not_due"   # 매뉴얼 점검은 요청할 때만
    store = env["store"]
    findings = store.list_findings(run_id=outcome["run_id"])
    assert {item["skill"] for item in findings} >= {SKILL_SPEC_SUMMARY, SKILL_E}
    # 사양 변경 분석 하나에서 [TC 점검]을 누르면 TC 비교 결과와 초안 Excel 이 생긴다 (REQ-QAINTEL-016).
    source = next(item for item in findings if item["skill"] == SKILL_SPEC_SUMMARY and item["subject"] == "VP-10")
    checked = _run(env, today=MONDAY.replace(hour=12), on_demand="tc-check", finding_id=source["id"])
    assert {item["skill"] for item in store.list_findings(run_id=checked["run_id"])} == {SKILL_COVERAGE}
    assert all(item["product"] == "vxvue" for item in findings)
    assert len(store.list_questions(unanswered_only=True)) == 1
    out_dir = env["cfg"].output_dir / checked["run_id"]
    workbook = load_workbook(out_dir / "impact_checklist_draft.xlsx")
    assert workbook.sheetnames == ["Checklist 초안", "Coverage", "Review"]
    assert [cell.value for cell in workbook["Checklist 초안"][1]][:4] == ["Category", "TC ID", "버전", "SRS No"]
    decisions = {row[4] for row in workbook["Coverage"].iter_rows(min_row=2, values_only=True)}
    assert {"UPDATE_EXISTING", "CREATE_NEW"} <= decisions
    audit = json.loads((out_dir / "audit.json").read_text(encoding="utf-8"))
    assert {task["task_id"] for task in audit["tasks"]} == {"COV-001"}
    assert (out_dir / "sent" / "COV-001.json").is_file()               # 보낸 입력이 남는다
    out_dir = env["cfg"].output_dir / outcome["run_id"]
    assert env["sent"] and "Finding" in env["sent"][0][0] and "VXvue" in env["sent"][0][0]
    assert "/qa-agent/runs/" in env["sent"][0][1]
    assert store.get_run(outcome["run_id"])["email_status"] == "sent"
    assert outcome["run_id"].endswith("-vxvue")
    # 새 스냅샷은 제품 폴더에 저장된다
    assert (env["cfg"].srs_snapshot_dir / "2026-09-28.json").is_file()


def test_rules_mismatch_skips_ai_but_keeps_deterministic_stages(env):
    runner = FakeRunner(_producer)
    outcome = _run(env, runner=runner, rules_state=RulesState(False, "1.12", "판이 다릅니다", None, None))
    assert runner.calls == []
    assert outcome["stages"]["SPEC_COVERAGE"]["status"] == "rules"
    assert outcome["stages"]["E"]["status"] == "ok"
    assert outcome["status"] == "PARTIAL" and "판이 다릅니다" in env["sent"][0][1]
    events = env["store"].list_events(product="vxvue", event_types=("SRS_UPDATED",))
    assert events and events[0]["analysis_status"] == "pending"          # 규칙이 고쳐지면 분석한다


def test_polarion_failure_does_not_stop_trace_gap(env):
    class Broken:
        def iter_workitems(self, query, **_):
            from app.modules.daily_qa.polarion import PolarionError
            raise PolarionError("HTTP 401")
            yield

    outcome = run_daily(env["cfg"], today=MONDAY, polarion_factory=lambda: Broken(), runner=FakeRunner(_producer),
                        inputs=env["inputs"], rules_state=rules_ok(env["tmp"]), send_email=env["sender"], store=env["store"])
    stages = outcome["stages"]
    assert stages["collect_srs"]["status"] == "failed" and "HTTP 401" in stages["collect_srs"]["note"]
    assert stages["collect_issues"]["status"] == "failed"
    assert stages["E"]["status"] == "ok"                              # 저장된 어제 스냅샷으로 돈다
    assert "스냅샷 사용" in stages["collect_srs"]["note"]
    assert outcome["status"] == "PARTIAL"


def test_empty_srs_result_is_not_saved_as_snapshot(env):
    env["polarion"].srs = []
    outcome = _run(env)
    assert outcome["stages"]["collect_srs"]["status"] == "failed"
    assert not (env["cfg"].srs_snapshot_dir / "2026-09-28.json").exists()


def test_broken_result_is_retried_once_then_failed(env):
    calls = []

    def bad(task, run):
        calls.append(task.task_id)
        (run.out_dir / f"{task.task_id}.json").write_text("not json", encoding="utf-8")

    outcome = _run(env, runner=FakeRunner(bad), today=TUESDAY)
    assert calls.count("COV-001") == 2
    assert outcome["stages"]["SPEC_COVERAGE"]["status"] == "failed"
    assert outcome["stages"]["E"]["status"] == "not_due"
    events = env["store"].list_events(product="vxvue", event_types=("SRS_UPDATED",))
    assert events[0]["analysis_status"] == "failed" and events[0]["attempts"] == 1


def test_rejected_findings_are_not_stored(env):
    def producer(task, run):
        payload = json.loads((run.in_dir / f"{task.task_id}.json").read_text(encoding="utf-8"))
        write_result(run, task, [{"subject": item["target"], "verdict": "NOT_COVERED", "summary": "근거 없음", "evidence": []}
                                 for item in payload.get("items", [])])

    outcome = _run(env, runner=FakeRunner(producer), today=TUESDAY)
    assert not env["store"].list_findings(run_id=outcome["run_id"], skill=SKILL_COVERAGE, analysis_type="SPEC_COVERAGE")
    assert "규칙 위반으로 버린 Finding" in outcome["stages"]["SPEC_COVERAGE"]["note"]


def test_dry_run_calls_no_ai_and_stores_nothing(env):
    outcome = _run(env, runner=None, dry_run=True, today=TUESDAY)
    assert outcome["stages"]["SPEC_COVERAGE"]["status"] == "skipped" and "dry-run" in outcome["stages"]["SPEC_COVERAGE"]["note"]
    assert env["store"].list_events(product="vxvue") == []
    assert not env["store"].list_findings(run_id=outcome["run_id"])
    assert not (env["cfg"].srs_snapshot_dir / "2026-09-29.json").exists()


def test_second_run_does_not_duplicate_trace_gap_findings(env):
    first = _run(env)
    second = _run(env, today=datetime(2026, 10, 5, 9, 0, tzinfo=timezone.utc))
    assert env["store"].list_findings(run_id=first["run_id"], skill=SKILL_E)
    assert not env["store"].list_findings(run_id=second["run_id"], skill=SKILL_E)


def test_trace_gap_dedupes_against_findings_saved_under_the_old_skill_name(env):
    """Skill 이름을 바꿔도(vxvue-trace-gap → qa-trace-gap) 옛 Finding 과 겹치지 않는다."""
    store = env["store"]
    store.create_run("old", False)
    store.add_finding("old", "vxvue-trace-gap", "E-weekly", {
        "subject": "VP-12", "verdict": "TC 없음", "summary": "옛 Finding",
        "evidence": [{"source_type": "srs", "location": "VP-12", "validity": "Current"}]})
    outcome = _run(env)
    subjects = [item["subject"] for item in store.list_findings(run_id=outcome["run_id"], skill=SKILL_E)]
    assert "VP-12" not in subjects


def test_concurrent_run_is_refused(env):
    lock = env["cfg"].lock_path
    lock.parent.mkdir(parents=True, exist_ok=True)
    lock.write_text("123", encoding="utf-8")
    with pytest.raises(RunLocked):
        _run(env)


def test_missing_token_skips_ai_stages(env):
    cfg = make_settings(env["cfg"].root, env["tmp"] / "ws", claude_token="")
    outcome = run_daily(cfg, today=TUESDAY, polarion_factory=lambda: env["polarion"], inputs=env["inputs"],
                        rules_state=rules_ok(env["tmp"]), send_email=env["sender"], store=env["store"])
    assert outcome["stages"]["SPEC_COVERAGE"]["status"] == "skipped"
    assert "CLAUDE_CODE_OAUTH_TOKEN" in outcome["stages"]["preflight"]["note"]


def test_manual_check_runs_only_when_requested(env):
    """주간 요일·변경이 있어도 자동으로는 돌지 않고, 요청하면 돈다 (REQ-QAINTEL-034)."""
    outcome = _run(env)
    assert outcome["stages"]["F"]["status"] == "not_due"
    requested = _run(env, today=MONDAY.replace(hour=12), on_demand="manual-check")
    assert requested["stages"]["F"]["status"] == "ok"
    assert any(task["task_id"] == "F-001" for task in json.loads(
        (env["cfg"].output_dir / requested["run_id"] / "audit.json").read_text(encoding="utf-8"))["tasks"])


def test_claude_tool_logs_are_collected_into_the_run_folder(env):
    """Validates: NFR-SEC-001 — 보낸 입력과 Claude 가 읽은 기록이 같은 실행 폴더에 모인다."""

    def producer(task, run):
        _producer(task, run)
        log_dir = run.run_dir / "logs"
        log_dir.mkdir(exist_ok=True)
        (log_dir / f"{task.task_id}.claude.json").write_text(
            json.dumps({"tool_calls": 1, "tool_call_log": [{"tool": "Read", "file_path": "runs/x/context/tc_index.jsonl"}]}),
            encoding="utf-8",
        )

    outcome = _run(env, runner=FakeRunner(producer), today=TUESDAY)
    logs = env["cfg"].output_dir / outcome["run_id"] / "claude_logs"
    saved = json.loads((logs / "COV-001.attempt1.claude.json").read_text(encoding="utf-8"))
    assert saved["tool_call_log"][0]["file_path"].endswith("tc_index.jsonl")
    assert (env["cfg"].output_dir / outcome["run_id"] / "sent" / "COV-001.json").is_file()


def test_scheduler_launches_detached_process_only_with_credentials(env, monkeypatch):
    from app.modules.daily_qa import scheduled_jobs

    launched = []

    class FakePopen:
        def __init__(self, args, **kwargs):
            launched.append((args, kwargs))
            self.pid = 4242

    no_polarion = PolarionSettings(host="", token="", project_id="VXvue", srs_query="type:srs", issue_query="type:issue")
    monkeypatch.setattr(scheduled_jobs, "load", lambda product=None: make_settings(env["cfg"].root, env["tmp"] / "ws", polarion=no_polarion))
    assert scheduled_jobs.launch_detached(popen=FakePopen)["status"] == "not_configured" and not launched

    monkeypatch.setattr(scheduled_jobs, "load", lambda product=None: env["cfg"])
    result = scheduled_jobs.launch_detached(trigger="manual", popen=FakePopen)
    assert result == {"status": "launched", "pid": 4242, "product": "vxvue"}
    args, kwargs = launched[0]
    assert args[1].endswith("run_daily_qa.py") and args[args.index("--product") + 1] == "VXvue"
    assert args[args.index("--trigger") + 1] == "manual"
    assert kwargs.get("start_new_session") or kwargs.get("creationflags")   # 웹 프로세스와 분리된다
