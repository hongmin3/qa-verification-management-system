"""Validates REQ-MANUAL-021 using synthetic manuals only."""
import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile, ZipInfo

import pytest
from openpyxl import load_workbook

from app.modules.manual_review.source_inventory import extract_inventory
from app.modules.manual_review.shared_review import ReviewStore, build_report

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def docx_bytes():
    body = f'''<w:document xmlns:w="{W}"><w:body>
      <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>1. Manual</w:t></w:r></w:p>
      <w:p><w:commentRangeStart w:id="7"/><w:r><w:t xml:space="preserve">  원문 A &amp; B </w:t><w:tab/><w:t>유지</w:t><w:br/><w:t>다음 줄</w:t></w:r><w:del><w:r><w:delText>옛 문구</w:delText></w:r></w:del><w:ins><w:r><w:t>새 문구</w:t></w:r></w:ins><w:commentRangeEnd w:id="7"/></w:p>
      <w:tbl><w:tr><w:tc><w:p><w:pPr><w:pStyle w:val="StrikeStyle"/></w:pPr><w:r><w:t>취소선 설명</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
      <w:p><w:moveFrom><w:r><w:t>출발</w:t></w:r></w:moveFrom><w:moveTo><w:r><w:t>도착</w:t></w:r></w:moveTo></w:p>
    </w:body></w:document>'''
    styles = f'''<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="Heading 1"/></w:style><w:style w:type="paragraph" w:styleId="StrikeBase"><w:rPr><w:strike/></w:rPr></w:style><w:style w:type="paragraph" w:styleId="StrikeStyle"><w:basedOn w:val="StrikeBase"/></w:style></w:styles>'''
    comments = f'''<w:comments xmlns:w="{W}"><w:comment w:id="7" w:author="QA"><w:p><w:r><w:t>Check literal wording.</w:t></w:r></w:p></w:comment></w:comments>'''
    stream = io.BytesIO()
    with ZipFile(stream, 'w') as z:
        z.writestr(ZipInfo('word/document.xml'), body)
        z.writestr(ZipInfo('word/styles.xml'), styles)
        z.writestr(ZipInfo('word/comments.xml'), comments)
    return stream.getvalue()


def store_source(tmp_path):
    store = ReviewStore(tmp_path)
    source = store.add(docx_bytes(), '테스트.docx')
    return store, source


def finding(identifier, paragraphs, **extra):
    return dict(id=identifier, title='검토 항목', decision='수정 필수', priority='높음',
                paragraphs=paragraphs, proposed='제안 문구', basis='합성 사양 1절',
                coverage='공통 절 대조', reason='조건 확인', **extra)


def test_inventory_preserves_literal_text_and_separates_deleted_history(tmp_path):
    path = tmp_path / 'manual.docx'
    path.write_bytes(docx_bytes())
    inventory = extract_inventory(path)
    assert len(inventory['paragraphs']) == 4
    p = inventory['paragraphs'][1]
    assert p['number'] == 2 and p['paragraph_index'] == 1
    assert p['text'] == '  원문 A & B \t유지\n다음 줄새 문구'
    assert p['before'] == '  원문 A & B \t유지\n다음 줄옛 문구'
    assert p['deleted'] == '옛 문구'
    assert p['section'] == '1. Manual'
    assert p['comment_ids'] == ['7']
    assert inventory['paragraphs'][2]['has_strike'] is True
    assert inventory['paragraphs'][3]['text'] == '도착'
    assert inventory['comments'][0]['paragraphs'] == [2]
    assert inventory['sha256'] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_report_uses_source_quotes_and_document_order_without_renumbering(tmp_path):
    store, source = store_source(tmp_path)
    rows = [finding('R01', [4]), finding('R03', [2])]
    result = build_report(store, source['source_id'], rows)
    payload = json.loads((Path(result['directory']) / 'report.json').read_text('utf-8'))
    assert [r['id'] for r in payload['findings']] == ['R03', 'R01']
    assert payload['findings'][0]['quotes'][0]['text'] == '  원문 A & B \t유지\n다음 줄새 문구'
    assert '옛 문구' not in payload['findings'][0]['current_text']
    html = (Path(result['directory']) / 'manual_review.html').read_text('utf-8')
    assert '  원문 A &amp; B \t유지\n다음 줄새 문구' in html
    workbook = load_workbook(Path(result['directory']) / 'manual_review.xlsx')
    ws = workbook['검토결과']
    headers = [c.value for c in ws[6]]
    assert ws.cell(7, headers.index('현재 문구 (원문 그대로)') + 1).value == payload['findings'][0]['current_text']
    assert workbook.sheetnames == ['검토결과', '메모검토', '릴리즈대조', '변경추적', '조사범위']
    assert store.inventory(source['source_id'])['sha256'] == source['source_id']


def test_report_rejects_fabricated_quote_unknown_paragraph_and_duplicate_id(tmp_path):
    store, source = store_source(tmp_path)
    invalid = [
        [finding('R01', [2], current_text='invented summary')],
        [finding('R01', [999])],
        [finding('R01', [2]), finding('R01', [4])],
    ]
    for rows in invalid:
        with pytest.raises(ValueError):
            build_report(store, source['source_id'], rows)


def test_source_id_and_hash_are_checked_and_upload_name_never_becomes_path(tmp_path):
    store = ReviewStore(tmp_path)
    source = store.add(docx_bytes(), '../../manual.docx')
    assert source['name'] == 'manual.docx'
    with pytest.raises(ValueError):
        store.inventory('../secrets')
    store.source_path(source['source_id']).write_bytes(b'tampered')
    with pytest.raises(ValueError, match='해시'):
        store.inventory(source['source_id'])


def test_pdf_page_mapping_does_not_choose_first_duplicate_match(tmp_path):
    import fitz
    store, source = store_source(tmp_path)
    pdf = fitz.open()
    for _ in range(2):
        page = pdf.new_page()
        page.insert_text((40, 40), '1. Manual')
    result = store.attach_pdf(source['source_id'], pdf.tobytes(), 'final')
    p = store.inventory(result['source_id'])['paragraphs'][0]
    assert p['pages']['final']['status'] == 'ambiguous'
    assert p['pages']['final']['pages'] == []


def test_pdf_source_retains_physical_page_and_requires_visual_revision_review(tmp_path):
    import fitz
    pdf = fitz.open()
    pdf.new_page().insert_text((40, 40), 'Exact PDF text: A + B')
    store = ReviewStore(tmp_path)
    source = store.add(pdf.tobytes(), 'manual.pdf')
    data = store.inventory(source['source_id'])
    assert data['paragraphs'][0]['physical_page'] == 1
    assert 'Exact PDF text: A + B' in data['paragraphs'][0]['text']
    assert data['revision_validity'] == 'visual_review_required'


def test_upload_page_and_generated_download_use_the_same_review_source(tmp_path):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.modules.manual_review.shared_routes import review_router
    app = FastAPI()
    app.include_router(review_router(lambda: tmp_path), prefix='/vxvue-desktop/manual-review')
    client = TestClient(app)
    assert client.get('/vxvue-desktop/manual-review').status_code == 200
    response = client.post('/vxvue-desktop/manual-review/upload', files={'manual': ('manual.docx', docx_bytes())})
    assert response.status_code == 200
    sid = hashlib.sha256(docx_bytes()).hexdigest()
    assert sid in response.text
    result = build_report(ReviewStore(tmp_path), sid, [finding('R01', [2])])
    response = client.get('/vxvue-desktop/manual-review/files/' + result['report_id'] + '/manual_review.html')
    assert response.status_code == 200
    assert '원문 그대로' in response.text
    assert client.get('/vxvue-desktop/manual-review/files/' + result['report_id'] + '/source.json').status_code == 404
    assert client.post('/vxvue-desktop/manual-review/upload', files={'manual': ('a.docx', b'not zip')}).status_code == 400
    assert client.post('/vxvue-desktop/manual-review/upload', headers={'Origin': 'https://other.invalid'}, files={'manual': ('a.docx', docx_bytes())}).status_code == 403


def test_desktop_definitions_and_calls_expose_shared_tools_without_bridge_change(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from app.modules.vxvue_chatbot import desktop
    monkeypatch.setattr(desktop, 'get_settings', lambda: SimpleNamespace(root=tmp_path))
    definitions = {t['name']: t for t in desktop.definitions()}
    assert 'get_manual_review' in definitions
    assert 'make_manual_review_report' in definitions
    assert definitions['get_manual_review']['annotations']['readOnlyHint'] is True
    assert definitions['make_manual_review_report']['annotations']['readOnlyHint'] is False
    _, source = store_source(tmp_path)
    result = desktop.call('get_manual_review', {'source_id': source['source_id'], 'view': 'paragraphs', 'paragraphs': [2]})
    assert result['results'][0]['text'] == '  원문 A & B \t유지\n다음 줄새 문구'
    report = desktop.call('make_manual_review_report', {'source_id': source['source_id'], 'findings': [finding('R03', [2])]})
    assert report['files'][0]['url'].startswith('/vxvue-desktop/manual-review/files/')
    assert 'directory' not in report


def test_special_characters_are_literal_and_unknown_font_symbols_block_quote(tmp_path):
    from app.modules.manual_review.source_inventory import _paragraph, W
    from xml.etree import ElementTree as ET
    p = ET.fromstring(f'<w:p xmlns:w="{W[1:-1]}"><w:r><w:t>A</w:t><w:noBreakHyphen/><w:t>B</w:t><w:softHyphen/><w:t>C</w:t></w:r></w:p>')
    assert _paragraph(p, {})['text'] == 'A\u2011B\u00adC'
    stream = io.BytesIO()
    with ZipFile(io.BytesIO(docx_bytes())) as original, ZipFile(stream, 'w') as target:
        for info in original.infolist():
            content = original.read(info.filename)
            if info.filename == 'word/document.xml':
                content = content.replace(b'<w:t>1. Manual</w:t>', b'<w:t>1. Manual</w:t><w:sym w:font="Wingdings" w:char="F020"/>')
            target.writestr(info, content)
    data = stream.getvalue()
    store = ReviewStore(tmp_path)
    source = store.add(data, 'symbols.docx')
    with pytest.raises(ValueError, match='특수'):
        build_report(store, source['source_id'], [finding('R01', [1])])


def test_single_and_double_strike_overrides_remain_independent():
    from app.modules.manual_review.source_inventory import _paragraph, W
    from xml.etree import ElementTree as ET
    styles = {'Single': {'name': '', 'base': '', 'strike': ET.fromstring(f'<w:strike xmlns:w="{W[1:-1]}"/>'), 'double': None}}
    para = ET.fromstring(f'<w:p xmlns:w="{W[1:-1]}"><w:pPr><w:pStyle w:val="Single"/></w:pPr><w:r><w:rPr><w:dstrike w:val="0"/></w:rPr><w:t>Visible strike</w:t></w:r></w:p>')
    assert _paragraph(para, styles)['has_strike'] is True
    para.find(W+'r/'+W+'rPr').append(ET.fromstring(f'<w:strike xmlns:w="{W[1:-1]}" w:val="0"/>'))
    assert _paragraph(para, styles)['has_strike'] is False


def test_pdf_spanning_pages_never_selects_earlier_duplicate():
    import fitz
    from app.modules.manual_review.shared_review import map_pdf
    pdf = fitz.open()
    for text in ('Repeated phrase', 'Unique previous anchor\nRepeated', 'phrase\nUnique following anchor'):
        pdf.new_page().insert_text((40, 40), text)
    inventory = {'paragraphs': [dict(number=i, text=t, raw=t, before=t, deleted='') for i, t in enumerate(
        ('Unique previous anchor', 'Repeated phrase', 'Unique following anchor'), 1)]}
    result = map_pdf(inventory, pdf.tobytes(), 'final')['paragraphs']['2']
    assert result['pages'] == [2, 3] or result['status'] in ('ambiguous', 'unverified')


def test_pdf_context_is_immutable_between_users(tmp_path):
    import fitz
    store, source = store_source(tmp_path)
    first = fitz.open()
    first.new_page().insert_text((40, 40), '1. Manual')
    second = fitz.open()
    second.new_page().insert_text((40, 40), 'Other cover')
    second.new_page().insert_text((40, 40), '1. Manual')
    a = store.attach_pdf(source['source_id'], first.tobytes(), 'final')
    b = store.attach_pdf(source['source_id'], second.tobytes(), 'final')
    assert a['source_id'] != b['source_id']
    assert store.inventory(a['source_id'])['paragraphs'][0]['pages']['final']['pages'] == [1]
    assert store.inventory(b['source_id'])['paragraphs'][0]['pages']['final']['pages'] == [2]
    assert store.inventory(source['source_id'])['pdfs'] == {}
    assert store.inventory(a['source_id'])['sha256'] == source['sha256']
    report = build_report(store, a['source_id'], [finding('R01', [1])])
    text = (Path(report['directory']) / 'manual_review.html').read_text('utf-8')
    assert '원본 SHA-256: ' + source['sha256'] in text


# Validates: REQ-MANUAL-021 (근거 문서 자체 검토 항목의 위치)
def test_source_finding_rejects_paragraph_numbers_of_another_document(tmp_path):
    store, source = store_source(tmp_path)
    sid = source['source_id']
    for numbers in ({'paragraphs': [1]}, {'paragraphs': [], 'quote_paragraphs': [1]}):
        item = dict(finding('SRC1', [], kind='source', location='릴리즈노트 문단 192'), **numbers)
        with pytest.raises(ValueError, match='근거 문서'):
            build_report(store, sid, [item])
    assert not (store.base / 'reports').exists()


def test_source_finding_requires_its_own_location_and_shows_no_manual_quote(tmp_path):
    store, source = store_source(tmp_path)
    sid = source['source_id']
    with pytest.raises(ValueError, match='location'):
        build_report(store, sid, [finding('SRC1', [], kind='source')])
    result = build_report(store, sid, [finding('SRC1', [], kind='source', location='릴리즈노트 문단 192')])
    row = json.loads((Path(result['directory']) / 'report.json').read_text('utf-8'))['findings'][0]
    assert row['quotes'] == [] and row['location'] == '릴리즈노트 문단 192' and row['pages_final'] == '해당 없음'
