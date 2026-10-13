"""등록 XLSX 사양서 전문·위치·캐시·검색.

Validates: REQ-PARSE-010, REQ-KNOW-006, REQ-KNOW-007, REQ-QAINTEL-011
"""
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.core import document_cache, knowledge_documents
from app.core.knowledge_documents import load_for_product
from app.core.storage import Storage
from app.modules.daily_qa.intelligence import Corpus, KnownIds
from app.modules.daily_qa.pipeline import default_spec_chunks
from app.parsers.document_parser import extract_document_text, parse_document
from tests.daily_qa_fixtures import make_settings


def workbook(path):
    book = Workbook()
    first = book.active
    first.title = "상태"
    first.append(["식별자", "설명"])
    first["A3"] = 0
    first["C3"] = False
    first["A43"] = "SYNTH-LATE"
    first["C43"] = "LateEntryMarker"
    second = book.create_sheet("설정")
    second["C4"] = "SecondSheetMarker"
    second["D4"] = "=SUM(A4:B4)"
    book.save(path)
    book.close()
    return path


def test_xlsx_text_reads_every_sheet_and_late_row_without_collapsing_empty_cells(tmp_path):
    path = workbook(tmp_path / "synthetic-spec.xlsx")
    before = path.read_bytes()
    text = extract_document_text(path)
    assert "상태 / 43행" in text and "C43: LateEntryMarker" in text
    assert "설정 / 4행" in text and "C4: SecondSheetMarker" in text
    assert "A3: 0" in text and "C3: False" in text
    assert "D4: =SUM(A4:B4)" in text
    assert "상태 / 2행" not in text
    assert path.read_bytes() == before


def test_xlsx_chunks_keep_sheet_row_locations_and_split_long_rows(tmp_path):
    path = workbook(tmp_path / "synthetic-spec.xlsx")
    book = load_workbook(path)
    book["설정"]["A8"] = "x" * 400 + "LONG_ROW_TAIL"
    book.save(path)
    book.close()
    chunks = parse_document(path, "synthetic", chunk_chars=100)
    assert all(len(chunk.text) <= 100 for chunk in chunks)
    assert len({chunk.chunk_id for chunk in chunks}) == len(chunks)
    assert any("상태" in chunk.heading and "43" in chunk.heading and "LateEntryMarker" in chunk.text for chunk in chunks)
    long = [chunk for chunk in chunks if "설정" in chunk.heading and "8" in chunk.heading]
    assert len(long) > 1 and "LONG_ROW_TAIL" in "".join(chunk.text for chunk in long)
    assert all(chunk.page == 0 for chunk in chunks)


def test_empty_xlsx_has_no_text_or_chunks(tmp_path):
    path = tmp_path / "empty.xlsx"
    book = Workbook()
    book.save(path)
    book.close()
    assert extract_document_text(path) == ""
    assert parse_document(path, "empty") == []


def test_broken_xlsx_is_reported_as_unreadable(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(ValueError, match="Excel 사양서를 읽을 수 없습니다"):
        parse_document(path, "broken")


def test_registered_xlsx_reuses_cache_and_reaches_daily_ai_search_with_original_location(tmp_path, monkeypatch):
    root = tmp_path / "root"
    (root / "data").mkdir(parents=True)
    cache = tmp_path / "cache"
    monkeypatch.setattr(document_cache, "get_settings", lambda: SimpleNamespace(path=lambda key: cache))
    path = workbook(tmp_path / "synthetic-spec.xlsx")
    cfg = make_settings(root, tmp_path / "workspace")
    storage = Storage(db_path=cfg.db_path)
    number = storage.add_document("specification", cfg.product, "1", "1", path.name, path)
    loaded = load_for_product(cfg.product, storage=storage, with_text=True, require_both=False)
    assert loaded.failures == [] and loaded.chunks
    assert "LateEntryMarker" in loaded.baseline_texts[0]
    assert (cache / f"{number}.json").is_file() and (cache / f"{number}.text").is_file()

    def no_reparse(*args, **kwargs):
        raise AssertionError("정상 저장본을 두고 원본을 다시 파싱함")
    monkeypatch.setattr(knowledge_documents, "parse_document", no_reparse)
    monkeypatch.setattr(knowledge_documents, "extract_document_text", no_reparse)
    second = load_for_product(cfg.product, storage=storage, with_text=True, require_both=False)
    assert second.failures == [] and second.chunks == loaded.chunks and second.baseline_texts == loaded.baseline_texts
    chunks, labels, unreadable = default_spec_chunks(cfg.product, cfg)
    assert unreadable == [] and chunks == loaded.chunks
    corpus = Corpus(cfg.product_profile, [], [], [], spec_chunks=chunks, chunk_labels=labels)
    candidates = corpus.spec_candidates(["LateEntryMarker"], "LateEntryMarker", 0, 3, KnownIds())
    evidence = next(item for item in candidates["spec_docs"] if "LateEntryMarker" in item["text"])
    assert path.name in evidence["location"] and "상태" in evidence["location"] and "43" in evidence["location"]
    assert "p." not in evidence["location"]
