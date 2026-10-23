"""자동 출처 사양서의 화면 등록 허용 · 지운 문서 기억 · 교체 파일 이름 맞추기.

Validates: REQ-KNOW-005, REQ-KNOW-018, REQ-KNOW-020, REQ-KNOW-022, REQ-KNOW-023

사양서가 `alm_crawler` 이고 `allow_manual: true` 인 가짜 제품으로 확인한다. DB·파싱 저장본·수집 기록·업로드
폴더는 모두 tmp 에 둔다. 실제 data/ 를 건드리지 않는다.
"""
from __future__ import annotations

import io
import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.core.product_config as product_config
import app.core.product_knowledge as product_knowledge
import app.modules.knowledge.router as knowledge_router
from app.core import document_cache, knowledge_status
from app.core.knowledge_registry import replacement_name
from app.core.storage import Storage
from app.main import app

ROOT = Path(__file__).resolve().parents[1]
ALLOWED_YAML = """
product: Auto Spec
version: "1.0"
specification:
  source: alm_crawler
  allow_manual: true
testcase:
  source: manual
qa_rules:
  required: false
"""
LOCKED_YAML = """
product: Locked Spec
version: "1.0"
specification:
  source: alm_crawler
testcase:
  source: manual
qa_rules:
  required: false
"""
VSS_OLD = "(사양서) VSS Specification (261006).xlsx"
VSS_NEW = "(사양서) VSS Specification (261023).xlsx"
ALM_PDF = "(사양서) Auto 사양서1(261023).pdf"


@pytest.fixture
def env(tmp_path, monkeypatch):
    products = tmp_path / "config" / "products"
    products.mkdir(parents=True)
    shutil.copy(ROOT / "config" / "products" / "vxvue.yaml", products / "vxvue.yaml")
    (products / "auto-spec.yaml").write_text(ALLOWED_YAML, encoding="utf-8")
    (products / "locked-spec.yaml").write_text(LOCKED_YAML, encoding="utf-8")
    monkeypatch.setattr(product_config, "_products_dir", lambda root=None: products)
    monkeypatch.setattr(product_knowledge, "get_settings", lambda: SimpleNamespace(root=tmp_path))
    monkeypatch.setattr(document_cache, "_cache_path", lambda document_id: tmp_path / "indexes" / f"{document_id}.json")
    monkeypatch.setattr(document_cache, "_text_cache_path", lambda document_id: tmp_path / "indexes" / f"{document_id}.text")
    uploads = {"storage.specification_dir": tmp_path / "specs", "storage.testcase_dir": tmp_path / "testcases"}
    for directory in uploads.values():
        directory.mkdir()
    monkeypatch.setattr(knowledge_router, "get_settings", lambda: SimpleNamespace(path=lambda key: uploads[key], root=tmp_path))
    storage = Storage(tmp_path / "app.db")
    monkeypatch.setattr(knowledge_router, "storage", storage)
    return SimpleNamespace(root=tmp_path, storage=storage, uploads=uploads, client=TestClient(app))


def _xlsx(text: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "VSS"
    sheet.append(["ID", "Requirement"])
    sheet.append(["VSS-001", text])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(env, slug: str, name: str, data: bytes, replace_id: int | None = None, kind: str = "specification"):
    form = {"kind": kind}
    if replace_id is not None:
        form["replace_id"] = str(replace_id)
    return env.client.post(f"/knowledge/products/{slug}/documents", data=form, files={"file": (name, data)}, follow_redirects=False)


def _write_collected(env, product: str, names: list[str]) -> None:
    """지식 폴더 수집 기록과 수집 사본. 등록은 `register_collected` 가 한다."""
    base = product_knowledge.product_dir(product, env.root)
    assets = []
    for name in names:
        path = base / "original" / "specification" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(name.encode("utf-8"))
        parsed = product_knowledge.parse_asset_name(name)
        assets.append({"kind": "specification", "file_name": name, "sha256": name, **parsed})
    (base / "manifest.json").write_text(json.dumps({"product": product, "assets": assets, "excluded": []}, ensure_ascii=False), encoding="utf-8")


def _register_collected(env, product: str = "Auto Spec"):
    return product_knowledge.register_collected(product, env.storage, env.root)


def _names(env, product: str = "Auto Spec") -> list[str]:
    return sorted(document["name"] for document in env.storage.active_documents("specification", product))


def _doc(env, name: str, product: str = "Auto Spec") -> dict:
    return next(document for document in env.storage.active_documents("specification", product) if document["name"] == name)


def _spec_asset(env, slug: str = "auto-spec"):
    health = knowledge_status.product_health(product_config.load_product_config(slug), env.storage, env.root)
    return next(asset for asset in health.assets if asset.kind == "specification")


def _today_kst() -> date:
    return datetime.now(ZoneInfo("Asia/Seoul")).date()


# -- 1. 화면 등록 허용 (REQ-KNOW-018, REQ-KNOW-020) -----------------------------------------


def test_allow_manual_opens_upload_replace_and_delete_in_an_alm_section(env):
    _write_collected(env, "Auto Spec", [VSS_OLD, ALM_PDF])
    _register_collected(env)

    asset = _spec_asset(env)

    assert asset.can_upload is True
    assert ".xlsx" in asset.accept.split(",")
    assert all(row.can_delete for row in asset.documents)
    assert asset.message.startswith("ALM 에서 자동으로 수집합니다.")
    assert asset.message.endswith("이 화면에서 새 문서를 등록하거나 현재 문서를 교체할 수도 있습니다.")
    assert "여기서 올릴 수 없습니다" not in asset.message

    page = env.client.get("/knowledge/products/auto-spec").text
    section = page[page.index('id="kind-specification"'):page.index('id="kind-testcase"')]
    assert "새 문서 등록" in section
    assert section.count(">교체<") == 2
    assert section.count(">삭제<") == 2
    assert 'accept=".pdf,.docx,.xlsx"' in section


def test_without_allow_manual_an_alm_section_stays_locked(env):
    response = _upload(env, "locked-spec", VSS_OLD, _xlsx("본문"))

    assert response.status_code == 409
    assert not any(env.uploads["storage.specification_dir"].iterdir())
    asset = _spec_asset(env, "locked-spec")
    assert asset.can_upload is False
    assert "여기서 올릴 수 없습니다" in asset.message


def test_vxvue_product_config_allows_manual_specifications():
    config = product_config.load_product_config("vxvue", ROOT)

    assert config.source_of("specification") == "alm_crawler"
    assert config.specification.allow_manual is True
    assert product_config.can_manual_upload(config, "specification") is True
    assert config.testcase.allow_manual is False


def test_unknown_source_is_never_opened_by_allow_manual():
    config = product_config.ProductConfig.model_validate(
        {"product": "X", "specification": {"source": "mystery", "allow_manual": True}})

    assert product_config.can_manual_upload(config, "specification") is False


# -- 2. 등록 (REQ-KNOW-018) ---------------------------------------------------------------


def test_sourceless_xlsx_specification_registers_and_is_searchable(env):
    response = _upload(env, "auto-spec", VSS_NEW, _xlsx("Detector battery status"))

    assert response.status_code == 303
    document = _doc(env, VSS_NEW)
    assert json.loads(document["metadata_json"])["source"] == "upload"
    assert "Detector battery status" in (document_cache.load_text(document["id"]) or "")


def test_alm_sync_registration_still_works_with_allow_manual(env):
    response = env.client.post("/knowledge/specification", data={"product": "Auto Spec", "source": "alm_crawler"},
                               files={"file": ("(사양서) Auto 사양서1(261023).docx", _docx("SRS 본문"))}, follow_redirects=False)

    assert response.status_code == 303
    assert _names(env) == ["(사양서) Auto 사양서1(261023).docx"]


def test_collected_specification_can_be_deleted_from_the_screen(env):
    _write_collected(env, "Auto Spec", [VSS_OLD, ALM_PDF])
    _register_collected(env)

    response = env.client.post(f"/knowledge/delete/{_doc(env, VSS_OLD)['id']}",
                               data={"next": "/knowledge/products/auto-spec"}, follow_redirects=False)

    assert response.status_code == 303
    assert response.headers["location"] == "/knowledge/products/auto-spec"
    assert _names(env) == [ALM_PDF]


# -- 3. 지운 문서 기억 (REQ-KNOW-022) ---------------------------------------------------------


def test_deleted_collected_document_is_not_registered_again(env):
    _write_collected(env, "Auto Spec", [VSS_OLD, ALM_PDF])
    _register_collected(env)
    env.client.post(f"/knowledge/delete/{_doc(env, VSS_OLD)['id']}", data={"next": "/knowledge"})

    # 다음 날 PC 가 같은 파일을 다시 보낸다 (서버 수집 사본도 다시 생긴다).
    _write_collected(env, "Auto Spec", [VSS_OLD, ALM_PDF])
    outcome = _register_collected(env)

    assert _names(env) == [ALM_PDF]
    assert "지운 문서라 건너뜀 1건" in outcome.detail
    assert env.storage.deleted_names("Auto Spec", "specification") == {VSS_OLD.casefold()}


def test_a_new_revision_of_a_deleted_document_is_registered(env):
    _write_collected(env, "Auto Spec", [VSS_OLD])
    _register_collected(env)
    env.client.post(f"/knowledge/delete/{_doc(env, VSS_OLD)['id']}", data={"next": "/knowledge"})

    _write_collected(env, "Auto Spec", [VSS_NEW])
    outcome = _register_collected(env)

    assert _names(env) == [VSS_NEW]
    assert "지운 문서라 건너뜀" not in outcome.detail


def test_deletion_by_the_automatic_source_is_not_remembered(env):
    env.client.post("/knowledge/specification", data={"product": "Auto Spec", "source": "alm_crawler"},
                    files={"file": ("(사양서) Auto 사양서1(261022).docx", _docx("SRS"))})
    document = _doc(env, "(사양서) Auto 사양서1(261022).docx")

    response = env.client.post(f"/knowledge/delete/{document['id']}", data={"source": "alm_crawler"}, follow_redirects=False)

    assert response.status_code == 303
    assert env.storage.deleted_names("Auto Spec", "specification") == set()


def test_uploading_the_same_name_again_clears_the_memory(env):
    _write_collected(env, "Auto Spec", [VSS_OLD])
    _register_collected(env)
    env.client.post(f"/knowledge/delete/{_doc(env, VSS_OLD)['id']}", data={"next": "/knowledge"})

    assert _upload(env, "auto-spec", VSS_OLD, _xlsx("다시 올림")).status_code == 303

    assert env.storage.deleted_names("Auto Spec", "specification") == set()
    assert _names(env) == [VSS_OLD]


def test_memory_is_kept_per_product_and_kind(env):
    env.storage.remember_deleted("Auto Spec", "specification", VSS_OLD)

    assert env.storage.deleted_names("Auto Spec", "testcase") == set()
    assert env.storage.deleted_names("Locked Spec", "specification") == set()
    env.storage.remember_deleted("Auto Spec", "specification", VSS_OLD.upper())   # 대소문자만 다른 이름은 같은 기록
    assert env.storage.deleted_names("Auto Spec", "specification") == {VSS_OLD.casefold()}


# -- 4. 교체 파일 이름 맞추기 (REQ-KNOW-023) ------------------------------------------------


@pytest.mark.parametrize(("target", "uploaded", "expected"), [
    ("(사양서) VSS Specification (261006).xlsx", "VSS_spec_최종.xlsx", "(사양서) VSS Specification (261023).xlsx"),
    ("(사양서) VSS Specification (261006).xlsx", "VSS(261020).xlsx", "(사양서) VSS Specification (261020).xlsx"),
    ("(매뉴얼) VXvue User Manual V1.0.11_KO.pdf", "manual_new_V1.0.12.pdf", "(매뉴얼) VXvue User Manual V1.0.12_KO.pdf"),
    ("(TC) RA16-148-002_VXvue_TestCase.xlsx", "tc.xlsx", "(TC) RA16-148-002_VXvue_TestCase.xlsx"),
    ("[QA 작성 규칙] VXvue_Rev1.17.md", "rules Rev1.18.md", "[QA 작성 규칙] VXvue_Rev1.18.md"),
    ("(사양서) VSS Specification (261006).xlsx", "VSS.pdf", "(사양서) VSS Specification (261023).pdf"),
])
def test_replacement_name_follows_the_target_name(target, uploaded, expected):
    assert replacement_name(target, uploaded, date(2026, 10, 23)) == expected


def test_replace_registers_the_upload_under_the_target_name_and_drops_the_target(env):
    _write_collected(env, "Auto Spec", [VSS_OLD, ALM_PDF])
    _register_collected(env)
    target = _doc(env, VSS_OLD)

    sent = _xlsx("새 VSS")
    response = _upload(env, "auto-spec", "VSS_spec_최종.xlsx", sent, replace_id=target["id"])

    expected = f"(사양서) VSS Specification ({_today_kst():%y%m%d}).xlsx"
    assert response.status_code == 303
    assert _names(env) == sorted([ALM_PDF, expected])
    download = env.client.get(f"/knowledge/download/{_doc(env, expected)['id']}")
    assert download.status_code == 200
    assert download.content == sent   # 내용은 그대로, 이름만 바뀐다
    # 다음 날 폴더에 옛 판이 그대로 있어도 화면에서 바꾼 판을 따른다.
    _write_collected(env, "Auto Spec", [VSS_OLD, ALM_PDF])
    _register_collected(env)
    assert _names(env) == sorted([ALM_PDF, expected])


def test_new_document_registration_keeps_the_uploaded_name(env):
    response = _upload(env, "auto-spec", "VSS_spec_최종.xlsx", _xlsx("새 VSS"))

    assert response.status_code == 303
    assert _names(env) == ["VSS_spec_최종.xlsx"]


def _tc_xlsx(tc_id: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TC"
    sheet.append(["TC ID", "Test Step", "Expected Result"])
    sheet.append([tc_id, "실행", "동작"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_replacing_a_testcase_also_follows_the_target_name(env):
    name = "(TC) RA16-148-002_AutoSpec_TestCase.xlsx"
    assert _upload(env, "auto-spec", name, _tc_xlsx("TC-001"), kind="testcase").status_code == 303
    target = next(d for d in env.storage.active_documents("testcase", "Auto Spec") if d["name"] == name)

    response = _upload(env, "auto-spec", "tc_new.xlsx", _tc_xlsx("TC-002"), replace_id=target["id"], kind="testcase")

    assert response.status_code == 303
    assert [d["name"] for d in env.storage.active_documents("testcase", "Auto Spec")] == [name]
