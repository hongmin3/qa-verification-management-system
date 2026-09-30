"""TEST-MANUAL-021: Claude 대화로 돌리는 매뉴얼 개정 검증 도구(scripts/manual_review_local.py)."""

import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app.modules.manual_review import local_review

ROOT = Path(__file__).resolve().parents[1]
W_NS_DECL = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
    "</Relationships>"
)


def _write_tracked_manual(path: Path) -> None:
    """변경 추적 표본: 기능 삽입 1건, 페이지 번호 삭제 1건(단순 변경), 그림 삽입 1건."""
    body = (
        "<w:p><w:r><w:t>4.2 로그인</w:t></w:r></w:p>"
        '<w:p><w:ins w:author="연구소" w:date="2026-08-01T00:00:00Z">'
        "<w:r><w:t>로그인을 5회 실패하면 계정이 10분 동안 잠긴다.</w:t></w:r></w:ins></w:p>"
        '<w:p><w:del w:author="연구소" w:date="2026-08-01T00:00:00Z">'
        "<w:r><w:delText>12</w:delText></w:r></w:del></w:p>"
        '<w:p><w:ins w:author="연구소" w:date="2026-08-01T00:00:00Z">'
        '<w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        '<wp:docPr id="1" name="그림 1" descr="로그인 잠금 화면"/></wp:inline></w:drawing></w:r></w:ins></w:p>'
    )
    xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W_NS_DECL}><w:body>{body}</w:body></w:document>'
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", xml)
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("_rels/.rels", _RELS)


def _write_knowledge_copy(root: Path) -> None:
    """지식 사본(data/product_knowledge/vxvue/)에 사양서 한 개를 둔다."""
    base = root / "data" / "product_knowledge" / "vxvue"
    spec_dir = base / "original" / "specification"
    spec_dir.mkdir(parents=True)
    spec = Document()
    spec.add_paragraph("SRS-LOGIN-003 로그인 잠금")
    spec.add_paragraph("사용자가 로그인을 5회 연속 실패하면 계정을 10분 동안 잠근다.")
    spec.save(str(spec_dir / "(사양서) VXvue 사양서1(260901).docx"))
    manifest = {"assets": [{"kind": "specification", "file_name": "(사양서) VXvue 사양서1(260901).docx", "sha256": "abc"}]}
    (base / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")


def _extract(tmp_path: Path, *extra: str) -> tuple[int, Path]:
    manual = tmp_path / "VXvue Service Manual.docx"
    if not manual.exists():
        _write_tracked_manual(manual)
    if not (tmp_path / "kroot").exists():
        _write_knowledge_copy(tmp_path / "kroot")
    code = local_review.main([
        "extract", "--manual", str(manual), "--product", "VXvue",
        "--knowledge-root", str(tmp_path / "kroot"), "--out-root", str(tmp_path / "out"), *extra,
    ])
    runs = sorted((tmp_path / "out").glob("*/changes.json")) if (tmp_path / "out").exists() else []
    return code, (runs[-1].parent if runs else tmp_path / "missing")


def _fill_decisions(run_dir: Path, decision_for_text=None) -> dict:
    changes = json.loads((run_dir / "changes.json").read_text(encoding="utf-8"))
    by_id = {item["change_id"]: item for item in changes["changes"]}
    decisions = json.loads((run_dir / "decisions.json").read_text(encoding="utf-8"))
    for row in decisions["decisions"]:
        change = by_id[row["change_id"]]
        candidates = [item["chunk_id"] for item in change["srs_candidates"]]
        if change["image_change"]:
            row.update(decision="PASS", confidence=0.9)
        else:
            row.update(
                decision="SUPPLEMENT_REQUIRED", confidence=0.7,
                evidence=[{"chunk_id": candidates[0], "type": "DIRECT_SPEC"}],
                problem="잠금 해제 방법이 빠졌습니다.", qa_comment="잠금 해제 방법(10분 뒤 자동 해제)을 추가해 주세요.",
            )
    (run_dir / "decisions.json").write_text(json.dumps(decisions, ensure_ascii=False, indent=2), encoding="utf-8")
    return decisions


# Validates: REQ-MANUAL-018
def test_extract_writes_changes_with_srs_candidates_and_decision_template(tmp_path):
    code, run_dir = _extract(tmp_path)

    assert code == 0
    changes = json.loads((run_dir / "changes.json").read_text(encoding="utf-8"))
    assert changes["summary"]["total_changes"] == 3
    assert changes["summary"]["functional_changes"] == 2
    login = next(item for item in changes["changes"] if "5회" in item["text"])
    assert login["functional"] is True
    assert login["srs_candidates"] and "5회 연속 실패" in login["srs_candidates"][0]["text"]
    image = next(item for item in changes["changes"] if item["image_change"])
    assert image["review_required"] is True
    page_number = next(item for item in changes["changes"] if item["text"] == "12")
    assert page_number["functional"] is False
    labels = {item["code"]: item["label"] for item in changes["judgment_values"]}
    assert labels["PASS"] == "문제없음" and len(labels) == 8

    template = json.loads((run_dir / "decisions.json").read_text(encoding="utf-8"))
    assert sorted(row["change_id"] for row in template["decisions"]) == sorted([login["change_id"], image["change_id"]])
    assert all(row["decision"] == "" for row in template["decisions"])


# Validates: REQ-MANUAL-018
def test_finish_writes_excel_comment_docx_and_comments_json(tmp_path):
    code, run_dir = _extract(tmp_path)
    assert code == 0
    _fill_decisions(run_dir)

    assert local_review.main(["finish", "--run-dir", str(run_dir)]) == 0

    workbook = load_workbook(run_dir / "manual_review.xlsx")
    rows = list(workbook["변경 판정"].iter_rows(values_only=True))
    header = rows[0]
    decision_col = header.index("판정")
    assert "보강 필요" in [row[decision_col] for row in rows[1:]]

    comment_docx = run_dir / "VXvue Service Manual_QA_Comment.docx"
    comments = list(Document(str(comment_docx)).comments)
    assert len(comments) == 2  # 보강 필요 1건 + 판정 불가로 바뀐 이미지 변경 1건
    assert comments[0].author == "VXvue QA AI"

    saved = json.loads((run_dir / "comments.json").read_text(encoding="utf-8"))
    assert len(saved["comments"]) == 2
    assert {item["status"] for item in saved["comments"]} == {"OPEN"}


# Validates: REQ-MANUAL-018, REQ-MANUAL-012
def test_finish_turns_image_pass_into_unable_to_determine(tmp_path):
    _, run_dir = _extract(tmp_path)
    _fill_decisions(run_dir)

    assert local_review.main(["finish", "--run-dir", str(run_dir)]) == 0

    final = json.loads((run_dir / "decisions.json").read_text(encoding="utf-8"))
    image_row = next(row for row in final["final"] if row["image_change"])
    assert image_row["decision"] == "UNABLE_TO_DETERMINE"
    assert image_row["ai_original_decision"] == "PASS"
    assert image_row["needs_human_review"] is True
    assert image_row["reason"]


# Validates: REQ-MANUAL-018
def test_finish_stops_with_exit_1_on_unknown_decision_value(tmp_path):
    _, run_dir = _extract(tmp_path)
    decisions = _fill_decisions(run_dir)
    decisions["decisions"][0]["decision"] = "LOOKS_FINE"
    (run_dir / "decisions.json").write_text(json.dumps(decisions, ensure_ascii=False), encoding="utf-8")

    assert local_review.main(["finish", "--run-dir", str(run_dir)]) == 1
    assert not (run_dir / "manual_review.xlsx").exists()


# Validates: REQ-MANUAL-018
def test_finish_stops_with_exit_1_when_evidence_is_not_a_sent_candidate(tmp_path):
    _, run_dir = _extract(tmp_path)
    decisions = _fill_decisions(run_dir)
    for row in decisions["decisions"]:
        if row["evidence"]:
            row["evidence"] = [{"chunk_id": "made-up-chunk", "type": "DIRECT_SPEC"}]
    (run_dir / "decisions.json").write_text(json.dumps(decisions, ensure_ascii=False), encoding="utf-8")

    assert local_review.main(["finish", "--run-dir", str(run_dir)]) == 1
    assert not (run_dir / "manual_review.xlsx").exists()


# Validates: REQ-MANUAL-018
def test_finish_stops_with_exit_1_when_a_decision_is_empty(tmp_path, capsys):
    _, run_dir = _extract(tmp_path)

    assert local_review.main(["finish", "--run-dir", str(run_dir)]) == 1
    assert "빈 판정" in capsys.readouterr().out


# Validates: REQ-MANUAL-018
def test_extract_stops_with_exit_2_for_missing_or_wrong_file(tmp_path):
    missing = local_review.main(["extract", "--manual", str(tmp_path / "없음.docx"), "--out-root", str(tmp_path / "out")])
    text_file = tmp_path / "manual.txt"
    text_file.write_text("x", encoding="utf-8")
    wrong = local_review.main(["extract", "--manual", str(text_file), "--out-root", str(tmp_path / "out")])

    assert missing == 2
    assert wrong == 2


# Validates: REQ-MANUAL-018
def test_extract_without_knowledge_copy_marks_every_change_for_human_review(tmp_path, capsys):
    manual = tmp_path / "VXvue Service Manual.docx"
    _write_tracked_manual(manual)

    code = local_review.main([
        "extract", "--manual", str(manual), "--product", "VXvue",
        "--knowledge-root", str(tmp_path / "empty"), "--out-root", str(tmp_path / "out"),
    ])

    assert code == 0
    run_dir = next((tmp_path / "out").glob("*/changes.json")).parent
    changes = json.loads((run_dir / "changes.json").read_text(encoding="utf-8"))
    assert changes["warnings"]
    assert all(item["needs_human_review"] for item in changes["changes"] if item["functional"])
    assert "사양서" in capsys.readouterr().out


# Validates: REQ-MANUAL-018, REQ-MANUAL-015
def test_prior_comments_are_carried_to_next_round(tmp_path):
    _, first_run = _extract(tmp_path)
    _fill_decisions(first_run)
    assert local_review.main(["finish", "--run-dir", str(first_run)]) == 0

    code, second_run = _extract(tmp_path, "--prior-comments", str(first_run / "comments.json"))

    assert code == 0
    changes = json.loads((second_run / "changes.json").read_text(encoding="utf-8"))
    assert changes["round_number"] == 1
    assert len(changes["prior_comments"]) == 2
    assert all("resolution_suggestion" in item for item in changes["prior_comments"])
    template = json.loads((second_run / "decisions.json").read_text(encoding="utf-8"))
    assert {row["status"] for row in template["prior_comments"]} == {"OPEN"}


# Validates: REQ-MANUAL-018
def test_command_line_entry_point_runs(tmp_path):
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "manual_review_local.py"), "extract", "--manual", str(tmp_path / "없음.docx"), "--out-root", str(tmp_path / "out")],
        capture_output=True, text=True, encoding="utf-8", cwd=str(ROOT),
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )

    assert result.returncode == 2
    assert "없음.docx" in result.stdout + result.stderr
