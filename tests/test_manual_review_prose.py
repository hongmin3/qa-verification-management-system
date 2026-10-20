"""REQ-MANUAL-022: preserve source quotes while rejecting unreviewed prose."""
import io
import json
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.modules.manual_review.shared_review import ReviewStore, build_report

BAD = '공통설정화면에서사용자가확인할사용설명을안내합니다.'
GOOD = '공통 설정 화면에서 사용자가 확인할 사용 설명을 안내합니다.'
ORIGINAL = '  원문안내문구의띄어쓰기는임의로변경하지않습니다.  '


def source(tmp_path):
    data = io.BytesIO()
    with ZipFile(data, 'w') as archive:
        archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t xml:space="preserve">' + ORIGINAL + '</w:t></w:r></w:p></w:body></w:document>')
    store = ReviewStore(tmp_path)
    return store, store.add(data.getvalue(), 'synthetic.docx')['source_id']


def finding(**changes):
    return dict(dict(id='R01', title='설명 보강', decision='보강 권장', priority='보통', paragraphs=[1],
                     proposed=GOOD, basis='합성 사양 1절', coverage='공통 절과 대조', reason='조건 확인', note='분석 설명'), **changes)


@pytest.mark.parametrize('field', ['title', 'proposed', 'basis', 'coverage', 'reason', 'note'])
def test_each_generated_field_requires_spacing_review_before_files(tmp_path, field):
    store, sid = source(tmp_path)
    with pytest.raises(ValueError, match='R01.*' + field):
        build_report(store, sid, [finding(**{field: BAD})])
    assert not (store.base / 'reports').exists()


@pytest.mark.parametrize('appendix', ['memo_review', 'release_review'])
def test_appendix_prose_is_checked_even_with_custom_column_names(tmp_path, appendix):
    store, sid = source(tmp_path)
    with pytest.raises(ValueError, match=appendix):
        build_report(store, sid, [finding()], **{appendix: [{'분석 설명': BAD}]})


def test_report_title_is_checked(tmp_path):
    store, sid = source(tmp_path)
    with pytest.raises(ValueError, match='title'):
        build_report(store, sid, [finding()], title=BAD)


def test_source_spacing_and_protected_technical_strings_are_preserved(tmp_path):
    store, sid = source(tmp_path)
    technical = 'VXvue/XIPL의 AutoRotation 확인. `한국어로된매우긴사용자정의설정식별자명`과 C:\\자료\\한국어로된매우긴사용자정의폴더이름\\guide.docx 및 https://example.invalid/한국어로된매우긴문서이름을확인하는경로'
    result = build_report(store, sid, [finding(note=technical)], memo_review=[{'메모 원문': ORIGINAL, '분석 설명': GOOD}])
    payload = json.loads((Path(result['directory']) / 'report.json').read_text('utf-8'))
    assert payload['findings'][0]['current_text'] == ORIGINAL
    assert payload['findings'][0]['note'] == technical
    assert payload['memo_review'][0]['메모 원문'] == ORIGINAL
    ws = load_workbook(Path(result['directory']) / 'manual_review.xlsx')['검토결과']
    assert ws['H7'].value == ORIGINAL and ws['I7'].value == GOOD and ws['N7'].value == technical


@pytest.mark.parametrize('path', [
    'docs/한국어로된매우긴사용자정의폴더이름/guide.docx',
    r'C:\manual files\한국어로된매우긴사용자정의폴더이름\guide.docx',
])
def test_relative_and_spaced_file_paths_are_preserved(tmp_path, path):
    store, sid = source(tmp_path)
    result = build_report(store, sid, [finding(note=path)])
    payload = json.loads((Path(result['directory']) / 'report.json').read_text('utf-8'))
    assert payload['findings'][0]['note'] == path
    ws = load_workbook(Path(result['directory']) / 'manual_review.xlsx')['검토결과']
    assert ws['N7'].value == path


def test_path_exemption_does_not_hide_adjacent_generated_prose(tmp_path):
    store, sid = source(tmp_path)
    note = r'C:\manual files\한국어로된매우긴사용자정의폴더이름\guide.docx ' + BAD
    with pytest.raises(ValueError, match='R01.note'):
        build_report(store, sid, [finding(note=note)])
    assert not (store.base / 'reports').exists()


def test_deep_extensionless_paths_do_not_trigger_nested_regex_backtracking():
    script = ('from app.modules.manual_review.prose_quality import spacing_candidates; '
              'assert spacing_candidates("C:/" + "folder/" * 40 + "file_without_extension") == []')
    result = subprocess.run([sys.executable, '-c', script], cwd=Path(__file__).resolve().parents[1],
                            capture_output=True, text=True, timeout=10)
    assert result.returncode == 0, result.stderr


def test_existing_connector_returns_correction_error_without_downloads(tmp_path, monkeypatch):
    from app.modules.vxvue_chatbot import desktop
    _, sid = source(tmp_path)
    monkeypatch.setattr(desktop, 'get_settings', lambda: SimpleNamespace(root=tmp_path))
    result = desktop.dispatch({'method': 'tools/call', 'params': {'name': 'make_manual_review_report',
                              'arguments': {'source_id': sid, 'findings': [finding(proposed=BAD)]}}})
    assert result['isError'] is True
    assert 'R01' in result['content'][0]['text'] and 'proposed' in result['content'][0]['text']
    assert '/files/' not in result['content'][0]['text']


def test_local_cli_uses_same_check(tmp_path, capsys):
    from app.modules.manual_review.shared_cli import main
    _, sid = source(tmp_path)
    path = tmp_path / 'findings.json'
    path.write_text(json.dumps([finding(proposed=BAD)], ensure_ascii=False), encoding='utf-8')
    assert main(['report', '--workspace-root', str(tmp_path), '--source-id', sid, '--findings', str(path)]) == 2
    assert 'proposed' in capsys.readouterr().out
