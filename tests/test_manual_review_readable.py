"""REQ-MANUAL-023: readable Korean analysis with unchanged literal evidence."""
import json
import html
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.modules.manual_review.shared_review import build_report
from app.modules.manual_review.shared_tools import definitions, call
from test_manual_review_shared import finding, store_source


def readable_finding():
    return finding('R01', [2], location='1.3.1 실행 방법',
        current_meaning='현재 문장은 실행 조건을 설명하고 있으나 경로가 빠져 있음.',
        evidence=[dict(source='합성 사양 VP-100', location='2절, 물리 3쪽',
                       quote='  Select Help > Manual.\nKeep UI labels.',
                       meaning='Help 메뉴에서 Manual을 선택하는 절차를 명시함.',
                       relation='매뉴얼의 실행 경로 설명을 보강할 근거임.')],
        short_comment='Help 메뉴의 Manual 실행 경로 추가 요청 드립니다.',
        language_context='intentional_korean_draft')


def test_current_quote_first_visible_and_copy_contains_only_short_memo(tmp_path):
    store, source = store_source(tmp_path)
    result = build_report(store, source['source_id'], [readable_finding()])
    folder = Path(result['directory'])
    payload = json.loads((folder / 'report.json').read_text('utf-8'))
    row = payload['findings'][0]
    document = (folder / 'manual_review.html').read_text('utf-8')
    card = re.search(r'<article id="R01"[^>]*>(.*?)</article>', document, re.S)[1]
    assert html.unescape(re.search(r'<pre>(.*?)</pre>', card, re.S)[1]) == row['current_text']
    assert '<details' not in card and ' hidden' not in card
    assert card.index(row['current_meaning']) > card.index('현재 문구')
    assert '대조 판단 (AI 분석)' in card and '번역 전 한국어 초안' in card
    assert html.unescape(re.search(r'<pre class="evidence-quote">(.*?)</pre>', card, re.S)[1]) == row['evidence'][0]['quote']
    memo = html.unescape(re.search(r'<pre class="review-comment">(.*?)</pre>\s*<button[^>]+data-copy="short-comment"', card, re.S)[1])
    assert memo == row['short_comment']
    assert '1.3.1' not in memo and row['basis'] not in memo
    assert row['decision'] == '수정 필수'  # Draft language never suppresses functional findings.
    ws = load_workbook(folder / 'manual_review.xlsx')['검토결과']
    assert ws['H7'].value == row['current_text']
    assert ws['I7'].value == row['proposed']
    headers = [c.value for c in ws[6]]
    assert ws.cell(7, headers.index('한국어 검토 코멘트') + 1).value == row['short_comment']
    assert row['evidence'][0]['quote'] in ws.cell(7, headers.index('근거 원문과 해설') + 1).value


def test_legacy_caller_gets_no_invented_translation_or_short_comment(tmp_path):
    store, source = store_source(tmp_path)
    result = build_report(store, source['source_id'], [finding('R01', [2])])
    document = (Path(result['directory']) / 'manual_review.html').read_text('utf-8')
    assert 'data-copy="short-comment"' not in document
    assert '짧은 메모 미작성' in document


@pytest.mark.parametrize('extra', [dict(short_comment=123), dict(current_meaning=['bad']),
    dict(evidence=[dict(source='invented')]), dict(language_context='automatic_defect'),
    dict(short_comment='조건을설명하는문장에는띄어쓰기가필요합니다')])
def test_invalid_explanations_fail_before_artifact_creation(tmp_path, extra):
    store, source = store_source(tmp_path)
    with pytest.raises(ValueError):
        build_report(store, source['source_id'], [finding('R01', [2], **extra)])
    assert not list((store.base / 'reports').glob('*'))


def test_mcp_contract_and_shared_call_support_new_and_legacy_fields(tmp_path):
    store, source = store_source(tmp_path)
    definition = next(x for x in definitions() if x['name'] == 'make_manual_review_report')
    props = definition['inputSchema']['properties']['findings']['items']['properties']
    assert {'current_meaning', 'evidence', 'short_comment', 'language_context'} <= props.keys()
    for row in [readable_finding(), finding('R02', [2])]:
        response = call('make_manual_review_report', dict(source_id=source['source_id'], findings=[row]), tmp_path, None)
        assert response['count'] == 1 and 'directory' not in response


def test_manual_prompts_preserve_korean_draft_and_require_grounded_korean_comments():
    from app.core.prompt_manager import load_prompt
    for name in ('manual_revision_quick', 'manual_revision_detail'):
        prompt = load_prompt(name)
        assert prompt.version == 4
        for phrase in ('한국어 초안', '기능·수치·사양 오류', 'UI·고유명·인용', '최신 통합 QA MD'):
            assert phrase in prompt.system_instruction
    detail = load_prompt('manual_revision_detail').system_instruction
    assert '한국어 뜻' in detail and '1~3문장' in detail and '수정 요청 드립니다' in detail
