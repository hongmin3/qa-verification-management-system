"""Validates: REQ-DAILY-004, REQ-DAILY-007 (TEST-DAILY-004)."""

from __future__ import annotations

import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.modules.daily_qa.checklist_xlsx import CHECKLIST_HEADERS, draft_rows_from_finding, write_draft
from app.modules.daily_qa.schema import SKILL_B, SKILL_C, ResultFileError, parse_result

EVIDENCE = [{"source_type": "srs", "location": "VP-1108 / 기능 요약", "summary": "", "validity": "Current"}]


def _write(path, findings, skill=SKILL_B, task="B-001"):
    path.write_text(json.dumps({"skill": skill, "task_id": task, "findings": findings}, ensure_ascii=False), encoding="utf-8")
    return path


def _finding(**overrides):
    base = {"subject": "VP-1108", "verdict": "수정 필수", "summary": "요약", "evidence": EVIDENCE}
    base.update(overrides)
    return base


def test_valid_finding_is_accepted(tmp_path):
    checked = parse_result(_write(tmp_path / "r.json", [_finding()]), SKILL_B, "B-001")
    assert len(checked.accepted) == 1 and not checked.rejected


@pytest.mark.parametrize(
    "overrides, reason",
    [
        ({"evidence": []}, "근거(evidence)가 없습니다"),
        ({"evidence": [{"source_type": "srs", "location": "사양서 어딘가", "validity": "Current"}]}, "문서 위치"),
        ({"verdict": "아마도 문제"}, "허용되지 않은 판정"),
        ({"action": "이슈를 Close 처리"}, "금지된 조치"),
        ({"action": "기존 TC 를 덮어쓰기"}, "금지된 조치"),
    ],
)
def test_rule_breaking_findings_are_rejected_with_reason(tmp_path, overrides, reason):
    checked = parse_result(_write(tmp_path / "r.json", [_finding(**overrides)]), SKILL_B, "B-001")
    assert not checked.accepted
    assert reason in checked.rejected[0]["reason"]


def test_tc_drafts_only_for_program_fixed(tmp_path):
    draft = {"kind": "수정확인", "title": "T", "test_step": "1. a", "expected_result": "1. b"}
    ok = _finding(verdict="신규 TC 필요", issue_type="Program Fixed", draft_tcs=[draft])
    bad = _finding(subject="VP-2", verdict="신규 TC 필요", issue_type="Inquiry", draft_tcs=[draft])
    checked = parse_result(_write(tmp_path / "r.json", [ok, bad], skill=SKILL_C, task="C-001"), SKILL_C, "C-001")
    assert [item.subject for item in checked.accepted] == ["VP-1108"]
    assert "Program Fixed" in checked.rejected[0]["reason"]


@pytest.mark.parametrize("content", [None, "not json", json.dumps({"skill": SKILL_B})])
def test_broken_result_file_raises_for_retry(tmp_path, content):
    path = tmp_path / "r.json"
    if content is not None:
        path.write_text(content, encoding="utf-8")
    with pytest.raises(ResultFileError):
        parse_result(path, SKILL_B, "B-001")


def test_result_for_another_task_is_refused(tmp_path):
    with pytest.raises(ResultFileError):
        parse_result(_write(tmp_path / "r.json", [], task="B-999"), SKILL_B, "B-001")


def _template(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "V1.1.0"
    sheet.append(list(CHECKLIST_HEADERS) + ["V1.1.0"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, name="맑은 고딕")
        cell.fill = PatternFill("solid", fgColor="FFD9E1F2")
    workbook.save(path)
    return path


def test_draft_workbook_matches_checklist_columns_and_separates_review(tmp_path):
    template = _template(tmp_path / "(TC) R-20-643_VXvue_변경사항 영향성평가_Checklist.xlsx")
    before = template.read_bytes()
    finding = {"subject": "VP-6669", "subject_title": "검색 조건", "draft_tcs": [
        {"kind": "수정확인", "srs_no": "VP-767", "title": "제목", "precondition": "1. p", "test_step": "1. s", "expected_result": "1. e"}]}
    rows = draft_rows_from_finding(finding, version="1.1.0")
    target = tmp_path / "draft.xlsx"
    info = write_draft(target, rows, [{"이슈 ID": "VP-6669", "판정": "신규 TC 필요", "요약": "AI 판단"}], template)
    workbook = load_workbook(target)
    sheet = workbook["Checklist 초안"]
    assert [cell.value for cell in sheet[1]] == list(CHECKLIST_HEADERS)
    assert sheet["D2"].value == "VP-767" and sheet["G2"].value == "제목"
    assert sheet["A1"].font.bold and sheet["A1"].fill.fgColor.rgb == "FFD9E1F2"
    assert sheet["A2"].fill.fill_type in (None, "none")           # 본문 셀에 배경색을 넣지 않는다
    assert "AI 판단" not in [cell.value for row in sheet.iter_rows() for cell in row]
    assert workbook["Review"]["E2"].value == "AI 판단"
    assert info["template_sheet"] == "V1.1.0"
    assert template.read_bytes() == before                         # 원본은 바뀌지 않는다
