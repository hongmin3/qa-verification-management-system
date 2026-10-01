"""사양 변경 Coverage 분석 (SPEC REQ-QAINTEL-016, REQ-QAINTEL-018).

Validates: REQ-QAINTEL-016, REQ-QAINTEL-018, REQ-QAINTEL-011
"""

from __future__ import annotations

from openpyxl import Workbook, load_workbook

from tests.daily_qa_fixtures import issue_item, srs_item
from tests.qa_intel_harness import DAY1, DAY2, DAY3, Harness

HEADERS = ["Category", "TC ID", "버전", "SRS No", "변경사항", "변경사항 상세", "Title", "Precondition", "Test Step", "Expected Result", "Test Data"]


def checklist(path, rows):
    """영향성평가 Checklist 모양의 TC 파일 (SRS 번호 열 + 머리글 서식)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1.1.0"
    sheet.append(["STC Category", "Old ID", "Old SRS ID", "Title", "Precondition", "Step Description", "Expected Reuslt"])
    for tc_id, srs_ref, title, expected in rows:
        sheet.append(["Changes Checklist", tc_id, srs_ref, title, "1. 로그인", "1. Retake 후 목록을 본다.", expected])
    template = workbook.create_sheet("Template")
    template.append(HEADERS)
    workbook.save(path)
    return path


def build(tmp_path, rows):
    tc = checklist(tmp_path / "(TC) R-20-643_VXvue_변경사항 영향성평가_Checklist.xlsx", rows)
    harness = Harness(tmp_path, tc_paths=[tc])
    harness.srs = [srs_item("VP-767", "07-01-01", "Retake 후 Protocol 목록", text="Retake 하면 목록을 다시 읽는다.")]
    harness.issues = [issue_item("VP-30301", "2026-09-01T00:00:00Z", ["VP-767"], review="lab_fixed", title="Retake 후 목록 안 바뀜"),
                      issue_item("VP-500", "2026-09-01T00:00:00Z", [], review="", title="무관한 로그인 이슈", description="로그인")]
    harness.run(DAY1)   # 기준
    return harness


def change_spec(harness, text="Retake 하면 목록을 다시 읽는다. Reload 뒤에도 목록 상태를 유지한다."):
    harness.srs[0] = srs_item("VP-767", "07-01-01", "Retake 후 Protocol 목록", text=text)


def draft(title, perspective="PERSISTENCE"):
    return {"kind": "Regression", "srs_no": "VP-767", "change": "Retake 후 목록", "title": title, "precondition": "1. 로그인",
            "test_step": "1. Retake 한다.\n2. Reload 한다.", "expected_result": "2. 목록 상태가 유지된다.",
            "perspective": perspective, "rationale": "Reload 뒤 상태 유지 검증 없음"}


def test_new_srs_without_tc_is_not_covered_and_gets_checklist_draft(tmp_path):
    harness = build(tmp_path, [])
    harness.srs.append(srs_item("VP-900", "09-01-01", "새 기능", text="새 버튼을 누르면 창이 열린다."))
    harness.script["VP-900"] = {"verdict": "NOT_COVERED", "draft_tcs": [{**draft("새 버튼 창 열기", "DIRECT_SPEC"), "srs_no": "VP-900"}],
                                "sections": {"change": {"summary": "새 기능"}, "checklist_coverage": {"tcs": []}}}
    outcome = harness.run(DAY2)
    finding = harness.store.list_findings(run_id=outcome["run_id"], analysis_type="SPEC_COVERAGE")[0]
    assert finding["verdict"] == "NOT_COVERED" and len(finding["draft_tcs"]) == 1
    workbook = load_workbook(harness.cfg.output_dir / outcome["run_id"] / "impact_checklist_draft.xlsx")
    rows = list(workbook["Checklist 초안"].iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS                                    # 제품 Checklist 형식 그대로
    assert rows[1][0] == "Changes Checklist" and rows[1][3] == "VP-900" and rows[1][6] == "새 버튼 창 열기"
    coverage = [row for row in workbook["Coverage"].iter_rows(min_row=2, values_only=True)]
    assert any(row[4] == "CREATE_NEW" for row in coverage)


def test_fully_covered_change_creates_no_new_tc_even_if_model_drafts_one(tmp_path):
    harness = build(tmp_path, [("TC-310", "VP-767", "Retake 후 목록", "1. Reload 뒤에도 목록 상태가 유지된다.")])
    change_spec(harness)
    harness.script["VP-767"] = {"verdict": "FULLY_COVERED", "draft_tcs": [draft("불필요")],
                                "sections": {"checklist_coverage": {"tcs": [{"tc_id": "TC-310", "decision": "KEEP", "reason": "Expected 까지 검증"}]}}}
    outcome = harness.run(DAY2)
    finding = harness.store.list_findings(run_id=outcome["run_id"], analysis_type="SPEC_COVERAGE")[0]
    assert finding["verdict"] == "FULLY_COVERED" and finding["draft_tcs"] == []
    assert not (harness.cfg.output_dir / outcome["run_id"] / "impact_checklist_draft.xlsx").exists() or \
        not list(load_workbook(harness.cfg.output_dir / outcome["run_id"] / "impact_checklist_draft.xlsx")["Checklist 초안"].iter_rows(min_row=2))


def test_old_expected_is_partially_covered_with_update_existing(tmp_path):
    harness = build(tmp_path, [("TC-310", "VP-767", "Retake 후 목록", "1. 목록이 바뀐다.")])
    change_spec(harness)
    harness.script["VP-767"] = {"verdict": "PARTIALLY_COVERED", "sections": {
        "checklist_coverage": {"tcs": [{"tc_id": "TC-310", "decision": "UPDATE_EXISTING", "reason": "Expected 가 옛 사양",
                                        "recommended_change": "Expected Result 1번 수정"}]},
        "alerts": [{"type": "TC_EXPECTED_OUTDATED", "detail": "Reload 조건 없음"}]}}
    outcome = harness.run(DAY2)
    finding = harness.store.list_findings(run_id=outcome["run_id"], analysis_type="SPEC_COVERAGE")[0]
    assert finding["sections"]["checklist_coverage"]["tcs"][0]["decision"] == "UPDATE_EXISTING"
    rows = list(load_workbook(harness.cfg.output_dir / outcome["run_id"] / "impact_checklist_draft.xlsx")["Coverage"].iter_rows(min_row=2, values_only=True))
    assert rows[0][4] == "UPDATE_EXISTING" and rows[0][6] == "Expected Result 1번 수정" and "TC-310" in rows[0][3]


def test_related_past_issue_without_tc_is_issue_covered_but_checklist_not_covered(tmp_path):
    harness = build(tmp_path, [])
    change_spec(harness)
    harness.script["VP-767"] = {"verdict": "NOT_COVERED", "sections": {
        "issue_coverage": {"issues": [{"issue_id": "VP-30301", "relation": "PAST_FIXED", "reason": "같은 Retake 동작"}]},
        "checklist_coverage": {"tcs": []}, "alerts": [{"type": "ISSUE_WITHOUT_TC"}]}}
    outcome = harness.run(DAY2)
    item = harness.payloads[-1]["items"][0]
    related = item["candidates"]["issues"]
    assert related[0]["id"] == "VP-30301" and related[0]["match"] == "연결 항목"   # 정확 연결이 먼저
    assert "과거 FIXED" in related[0]["relation_hint"]
    finding = harness.store.list_findings(run_id=outcome["run_id"], analysis_type="SPEC_COVERAGE")[0]
    assert finding["sections"]["issue_coverage"]["has_related"] and finding["verdict"] == "NOT_COVERED"
    assert finding["related_ids"]["issues"] == ["VP-30301"]


def test_spec_review_required_holds_tc_generation(tmp_path):
    harness = build(tmp_path, [])
    change_spec(harness, "Retake 하면 경우에 따라 목록을 다시 읽을 수 있다.")
    harness.script["VP-767"] = {"verdict": "SPEC_REVIEW_REQUIRED", "draft_tcs": [draft("근거 없는 TC")], "evidence": [],
                                "sections": {"tc_hold": {"reason": "Expected 근거 부족", "questions": ["어떤 경우인가?"]}}}
    outcome = harness.run(DAY2)
    finding = harness.store.list_findings(run_id=outcome["run_id"], analysis_type="SPEC_COVERAGE")[0]
    assert finding["draft_tcs"] == [] and finding["confidence"] == "Review Needed"
    rows = list(load_workbook(harness.cfg.output_dir / outcome["run_id"] / "impact_checklist_draft.xlsx")["Coverage"].iter_rows(min_row=2, values_only=True))
    assert rows[0][4] == "SPEC_REVIEW_REQUIRED"


def test_hallucinated_issue_and_tc_ids_are_removed(tmp_path):
    harness = build(tmp_path, [("TC-310", "VP-767", "Retake 후 목록", "1. 목록이 바뀐다.")])
    change_spec(harness)
    harness.script["VP-767"] = {"verdict": "PARTIALLY_COVERED", "sections": {
        "issue_coverage": {"issues": [{"issue_id": "VP-99999", "relation": "PAST_FIXED"}, {"issue_id": "VP-30301", "relation": "PAST_FIXED"}]},
        "checklist_coverage": {"tcs": [{"tc_id": "TC-DOES-NOT-EXIST", "decision": "UPDATE_EXISTING"},
                                       {"tc_id": "TC-310", "decision": "UPDATE_EXISTING"}]}}}
    outcome = harness.run(DAY2)
    finding = harness.store.list_findings(run_id=outcome["run_id"], analysis_type="SPEC_COVERAGE")[0]
    assert [entry["issue_id"] for entry in finding["sections"]["issue_coverage"]["issues"]] == ["VP-30301"]
    assert [entry["tc_id"] for entry in finding["sections"]["checklist_coverage"]["tcs"]] == ["TC-310"]
    removed = {item["value"] for item in finding["sections"]["validation"]["removed"]}
    assert {"VP-99999", "TC-DOES-NOT-EXIST"} <= removed


def test_same_change_rerun_does_not_create_duplicate_drafts(tmp_path):
    harness = build(tmp_path, [])
    change_spec(harness)
    harness.script["VP-767"] = {"verdict": "NOT_COVERED", "draft_tcs": [draft("Reload 뒤 상태 유지")]}
    harness.run(DAY2)
    again = harness.run(DAY2.replace(hour=12))   # 같은 날 다시 실행: 어제와 비교하면 같은 변경이다
    assert harness.skills_called().count("qa-spec-coverage-analysis") == 1
    assert harness.store.list_findings(run_id=again["run_id"], analysis_type="SPEC_COVERAGE") == []
    assert len(harness.events(event_types=("SRS_UPDATED",))) == 1


def test_no_change_day_calls_no_coverage_analyzer(tmp_path):
    harness = build(tmp_path, [])
    outcome = harness.run(DAY2)
    assert outcome["status"] == "NO_CHANGE" and harness.calls == []


def test_tc_candidates_follow_the_search_order_and_checklist_first(tmp_path):
    harness = build(tmp_path, [
        ("TC-LEGACY", "07-01-01", "옛 번호로 연결", "1. x"),
        ("TC-DIRECT", "VP-767", "새 번호로 연결", "1. x"),
        ("TC-ISSUE", "VP-1", "VP-30301 재발 확인", "1. x"),
    ])
    change_spec(harness)
    harness.script["VP-767"] = {"verdict": "PARTIALLY_COVERED"}
    harness.run(DAY2)
    tcs = harness.payloads[-1]["items"][0]["candidates"]["tcs"]
    assert [tc["tc_id"] for tc in tcs[:3]] == ["TC-DIRECT", "TC-LEGACY", "TC-ISSUE"]
    assert [tc["match"] for tc in tcs[:3]] == ["SRS 번호로 연결", "Legacy 번호로 연결", "관련 이슈 번호가 적힘"]
    assert all(tc["is_checklist"] for tc in tcs[:3])


def test_change_payload_carries_sentence_level_diff(tmp_path):
    harness = build(tmp_path, [])
    change_spec(harness)
    harness.run(DAY3)
    change = harness.payloads[-1]["items"][0]["change"]
    assert change["after"]["added_sentences"] == ["Reload 뒤에도 목록 상태를 유지한다."]
    assert change["events"] == ["SRS_UPDATED"]
