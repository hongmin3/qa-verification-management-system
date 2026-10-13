from __future__ import annotations

import zipfile
from collections.abc import Iterator
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.core.document_schemas import SpecificationChunk
from app.parsers.pdf_parser import extract_pdf_text, parse_specification

WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def extract_docx_text(path: Path) -> str:
    """Extract paragraphs and table cell text from a DOCX in document order."""
    try:
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        lines: list[str] = []
        for paragraph in root.iter(f"{WORD_NS}p"):
            text = "".join(node.text or "" for node in paragraph.iter(f"{WORD_NS}t")).strip()
            if text:
                lines.append(text)
        return "\n".join(lines)
    except Exception as exc:
        raise ValueError(f"Word 문서를 읽을 수 없습니다: {path.name}") from exc


def extract_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix == ".docx":
        return extract_docx_text(path)
    if suffix == ".xlsx":
        return "\n".join(text for _, _, _, text in _xlsx_rows(path))
    raise ValueError(f"지원하지 않는 문서 형식입니다: {suffix}")


def _xlsx_rows(path: Path) -> Iterator[tuple[int, str, int, str]]:
    """모든 시트의 값을 원래 행·셀 주소와 함께 읽는다(REQ-PARSE-010)."""
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        for sheet_number, sheet in enumerate(workbook.worksheets, start=1):
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = []
                for column, value in enumerate(values, start=1):
                    if value is None:
                        continue
                    text = str(value).strip()
                    if text:
                        cells.append(f"{get_column_letter(column)}{row_number}: {text}")
                if cells:
                    yield sheet_number, sheet.title, row_number, f"{sheet.title} / {row_number}행 | {' | '.join(cells)}"
    except Exception as exc:
        raise ValueError(f"Excel 사양서를 읽을 수 없습니다: {path.name}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _parse_xlsx(path: Path, document_id: str, chunk_chars: int) -> list[SpecificationChunk]:
    if chunk_chars <= 0:
        raise ValueError("조각 크기는 0보다 커야 합니다.")
    chunks = []
    sheet_number, sheet_name, sequence = 0, "", 0
    first_row, last_row, size = 0, 0, 0
    lines: list[str] = []

    def flush() -> None:
        nonlocal lines, size, sequence
        if not lines:
            return
        rows = f"{first_row}행" if first_row == last_row else f"{first_row}~{last_row}행"
        chunks.append(SpecificationChunk(chunk_id=f"{document_id}-s{sheet_number}-{sequence}",
                                         document_id=document_id, page=0,
                                         heading=f"{sheet_name} / {rows}", text="\n".join(lines)))
        sequence += 1
        lines, size = [], 0

    for current_sheet, title, row, text in _xlsx_rows(path):
        if current_sheet != sheet_number:
            flush()
            sheet_number, sheet_name, sequence = current_sheet, title, 0
        for offset in range(0, len(text), chunk_chars):
            part = text[offset:offset + chunk_chars]
            if lines and size + 1 + len(part) > chunk_chars:
                flush()
            if not lines:
                first_row = row
            last_row = row
            size += len(part) + (1 if lines else 0)
            lines.append(part)
    flush()
    return chunks


def parse_document(path: Path, document_id: str, chunk_chars: int = 1800) -> list[SpecificationChunk]:
    if path.suffix.lower() == ".pdf":
        return parse_specification(path, document_id, chunk_chars)
    if path.suffix.lower() == ".xlsx":
        return _parse_xlsx(path, document_id, chunk_chars)
    text = extract_document_text(path).strip()
    if not text:
        return []
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    heading = lines[0][:160] if lines else path.stem
    return [
        SpecificationChunk(
            chunk_id=f"{document_id}-p1-{offset // chunk_chars}",
            document_id=document_id,
            page=1,
            heading=heading,
            text=text[offset : offset + chunk_chars].strip(),
        )
        for offset in range(0, len(text), chunk_chars)
    ]
