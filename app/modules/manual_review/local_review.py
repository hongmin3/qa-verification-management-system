"""Claude 대화로 돌리는 매뉴얼 개정 검증 (REQ-MANUAL-018).

화면 검증(`reviewer.py`)과 같은 추출·거르기·근거 찾기·Release 대조·Comment 쓰기 코드를 쓰지만
AI 판정은 사용자가 연 Claude 세션이 한다. 이 도구는 외부 AI를 부르지 않고 핵심 앱 DB에도
쓰지 않는다. 두 단계로 나뉜다.

    extract  변경을 뽑아 changes.json 과 판정 기록 틀 decisions.json 을 쓴다.
    finish   decisions.json 을 검사하고 결과 Excel·Comment DOCX·comments.json 을 만든다.

명령줄 입구는 `scripts/manual_review_local.py` 다. 종료 코드: 0 정상, 1 판정 기록 오류,
2 입력 파일 문제.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from app.core.document_schemas import SpecificationChunk
from app.modules.manual_review.change_filter import is_functional_change
from app.modules.manual_review.comment_resolution import suggest_prior_comments
from app.modules.manual_review.comment_writer import insert_comments_detailed
from app.modules.manual_review.docx_track_changes import TrackedChange, extract_track_changes
from app.modules.manual_review.pdf_revision_diff import extract_pdf_revision_diff
from app.modules.manual_review.release_scope import extract_design_review_changes, extract_release_note_changes, match_release_changes
from app.modules.manual_review.review_signals import apply_human_review_signals, is_image_change
from app.modules.manual_review.schemas import (
    JUDGMENT_LABELS_KO,
    CommentStatus,
    Evidence,
    EvidenceType,
    ManualChangeJudgment,
    ManualJudgment,
)
from app.modules.manual_review.srs_evidence import search_candidates

TOOL_NAME = "manual_review_local"
FORMAT_VERSION = 1
DEFAULT_MAX_CANDIDATES = 6
SRS_NOT_AVAILABLE_REASON = "SRS_NOT_AVAILABLE"
OPEN_COMMENT_STATUSES = {CommentStatus.OPEN.value, CommentStatus.NOT_RESOLVED.value, CommentStatus.REOPENED.value}

EXIT_OK = 0
EXIT_DECISION_ERROR = 1
EXIT_INPUT_ERROR = 2

KIND_LABELS = {
    "insertion": "추가",
    "deletion": "삭제",
    "move_from": "이동(출발)",
    "move_to": "이동(도착)",
    "pdf_addition": "PDF 추가",
    "pdf_deletion": "PDF 삭제",
    "pdf_modification": "PDF 수정",
    "pdf_image_change": "PDF 이미지 변경",
}
RELEASE_STATUS_LABELS = {"FOUND": "찾음", "MISSING_SUSPECTED": "누락 의심"}
COMMENT_STATUS_LABELS = {
    "OPEN": "처리 전",
    "RESOLVED": "해결됨",
    "NOT_RESOLVED": "미해결",
    "REOPENED": "재오픈",
    "IGNORED_BY_QA": "QA 판단으로 제외",
}


class InputError(Exception):
    """입력 파일이 없거나 형식이 다르다 (종료 코드 2)."""


def kind_label(kind: str) -> str:
    if kind.startswith("image_"):
        return f"이미지 {KIND_LABELS.get(kind.removeprefix('image_'), kind)}"
    return KIND_LABELS.get(kind, kind)


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _write_json(path: Path, value) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _read_json(path: Path) -> dict:
    if not path.is_file():
        raise InputError(f"파일이 없습니다: {path}")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise InputError(f"JSON 파일을 읽을 수 없습니다: {path} ({exc})") from exc


def _require_file(path: Path | None, label: str, suffixes: tuple[str, ...]) -> Path:
    if path is None:
        raise InputError(f"{label} 경로가 없습니다.")
    if not path.is_file():
        raise InputError(f"{label} 파일이 없습니다: {path}")
    if path.suffix.lower() not in suffixes:
        raise InputError(f"{label} 형식이 다릅니다 ({' / '.join(suffixes)} 만 받습니다): {path}")
    return path


# ---------------------------------------------------------------- 사양서 (지식 사본)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def load_knowledge_chunks(product: str, knowledge_root: Path | None, cache_dir: Path) -> tuple[list[SpecificationChunk], list[str], list[str]]:
    """지식 사본(`data/product_knowledge/<제품>/`)의 사양서를 조각으로 나눈다.

    한 번 나눈 결과는 파일 내용 해시별로 cache_dir 에 두고 다음 실행에서 다시 쓴다.
    반환값: (조각, 쓴 사양서 이름, 경고)."""
    from app.core.product_knowledge import KIND_SPECIFICATION, collected_assets, collected_path
    from app.parsers.document_parser import parse_document

    chunks: list[SpecificationChunk] = []
    names: list[str] = []
    warnings: list[str] = []
    for asset in collected_assets(product, KIND_SPECIFICATION, knowledge_root):
        path = collected_path(product, asset, knowledge_root)
        if not path.is_file():
            warnings.append(f"지식 사본에 사양서 파일이 없어 건너뜁니다: {asset.get('file_name')}")
            continue
        cache_path = cache_dir / f"{_file_sha256(path)}.json"
        parsed: list[SpecificationChunk] | None = None
        if cache_path.is_file():
            try:
                parsed = [SpecificationChunk.model_validate(item) for item in json.loads(cache_path.read_text(encoding="utf-8"))]
            except (OSError, ValueError):
                parsed = None
        if parsed is None:
            try:
                parsed = parse_document(path, path.stem)
            except Exception as exc:  # 사양서 한 개를 못 읽어도 나머지로 계속한다
                warnings.append(f"사양서를 읽지 못해 건너뜁니다: {path.name} ({exc})")
                continue
            try:
                cache_dir.mkdir(parents=True, exist_ok=True)
                cache_path.write_text(json.dumps([item.model_dump(mode="json") for item in parsed], ensure_ascii=False), encoding="utf-8")
            except OSError:
                pass  # 캐시는 속도만 높인다
        chunks.extend(parsed)
        names.append(path.name)
    if not chunks:
        warnings.append("지식 사본에 사양서가 없어 근거 없이 진행합니다. 모든 변경에 사람 검토 필요 표시를 켭니다.")
    return chunks, names, warnings


def _candidate_record(chunk: SpecificationChunk) -> dict:
    return {
        "chunk_id": chunk.chunk_id,
        "document": chunk.document_id,
        "page": chunk.page,
        "heading": chunk.heading,
        "text": chunk.text,
    }


# ---------------------------------------------------------------- extract


def _extract_changes(manual: Path, previous: Path | None) -> list[TrackedChange]:
    try:
        if manual.suffix.lower() == ".pdf":
            if previous is None:
                raise InputError("PDF 매뉴얼은 --previous 로 이전 판 PDF 가 필요합니다.")
            return extract_pdf_revision_diff(previous, manual).changes
        return extract_track_changes(manual).changes
    except InputError:
        raise
    except Exception as exc:
        raise InputError(f"매뉴얼을 읽을 수 없습니다: {manual} ({exc})") from exc


def _release_items(release_notes: list[Path], design_reviews: list[Path]) -> list[tuple[str, object]]:
    from app.parsers.document_parser import extract_document_text

    items: list[tuple[str, object]] = []
    for path in release_notes:
        try:
            items.extend(("release_note", item) for item in extract_release_note_changes(extract_document_text(path), path.name))
        except Exception as exc:
            raise InputError(f"Release Note 를 읽을 수 없습니다: {path} ({exc})") from exc
    for path in design_reviews:
        try:
            items.extend(("design_review", item) for item in extract_design_review_changes(extract_document_text(path), path.name))
        except Exception as exc:
            raise InputError(f"설계검토보고서를 읽을 수 없습니다: {path} ({exc})") from exc
    return items


def _load_prior_comments(path: Path | None) -> tuple[int, list[dict]]:
    if path is None:
        return 0, []
    data = _read_json(path)
    if data.get("tool") != TOOL_NAME or not isinstance(data.get("comments"), list):
        raise InputError(f"이 도구가 만든 comments.json 이 아닙니다: {path}")
    return int(data.get("round_number", 0)) + 1, list(data["comments"])


def _unique_run_dir(out_root: Path, manual: Path) -> Path:
    base = f"{manual.stem}_{_now().astimezone().strftime('%Y%m%d-%H%M%S')}"
    candidate = out_root / base
    index = 2
    while candidate.exists():
        candidate = out_root / f"{base}-{index}"
        index += 1
    candidate.mkdir(parents=True)
    return candidate


def run_extract(args: argparse.Namespace) -> int:
    manual = _require_file(Path(args.manual), "개정 매뉴얼", (".docx", ".pdf"))
    previous = _require_file(Path(args.previous), "이전 판 매뉴얼", (".pdf",)) if args.previous else None
    release_notes = [_require_file(Path(item), "Release Note", (".docx", ".pdf")) for item in args.release_note]
    design_reviews = [_require_file(Path(item), "설계검토보고서", (".docx", ".pdf")) for item in args.design_review]
    prior_path = Path(args.prior_comments) if args.prior_comments else None
    round_number, prior_comments = _load_prior_comments(prior_path)

    tracked = _extract_changes(manual, previous)
    release_items = _release_items(release_notes, design_reviews)
    out_root = Path(args.out_root) if args.out_root else _repo_root() / "output" / "manual_review_local"
    knowledge_root = Path(args.knowledge_root) if args.knowledge_root else None
    chunks, spec_names, warnings = load_knowledge_chunks(args.product, knowledge_root, out_root / ".cache" / "srs")
    no_srs = not chunks

    changes: list[dict] = []
    for index, change in enumerate(tracked, start=1):
        changes.append({
            "change_id": f"C{index:03d}",
            "kind": change.kind,
            "kind_label": kind_label(change.kind),
            "author": change.author,
            "date": change.date,
            "paragraph_index": change.paragraph_index,
            "change_index_in_paragraph": change.change_index_in_paragraph,
            "source_page": change.source_page,
            "text": change.text,
            "functional": is_functional_change(change),
            "review_required": change.review_required,
            "image_change": is_image_change(change.kind),
            "needs_human_review": change.review_required or no_srs,
            "srs_candidates": [],
            "release_context": [],
        })
    functional = [item for item in changes if item["functional"]]
    for item in functional:
        item["srs_candidates"] = [_candidate_record(chunk) for chunk in search_candidates(chunks, item["text"], args.max_candidates)]

    release_findings: list[dict] = []
    if release_items:
        matched = match_release_changes([item for _, item in release_items], [(item["change_id"], item["text"]) for item in functional])
        by_id = {item["change_id"]: item for item in functional}
        for (source, release_change), (_, matched_id) in zip(release_items, matched):
            finding = {
                "source": source,
                "source_document": release_change.source_document,
                "category": release_change.category,
                "title": release_change.title,
                "description": release_change.description,
                "result_status": release_change.result_status,
                "status": "FOUND" if matched_id else "MISSING_SUSPECTED",
                "matched_change_id": matched_id,
            }
            release_findings.append(finding)
            if matched_id:
                by_id[matched_id]["release_context"].append({key: finding[key] for key in ("source", "category", "title", "description", "result_status")})
                if release_change.result_status == "FAIL":
                    by_id[matched_id]["needs_human_review"] = True

    open_prior = [item for item in prior_comments if item.get("status") in OPEN_COMMENT_STATUSES]
    closed_prior = [item for item in prior_comments if item.get("status") not in OPEN_COMMENT_STATUSES]
    current_for_suggestion = [{"id": item["change_id"], "functional": item["functional"], "text": item["text"], "kind": item["kind"]} for item in changes]
    open_prior = suggest_prior_comments(open_prior, current_for_suggestion)

    missing = sum(1 for item in release_findings if item["status"] == "MISSING_SUSPECTED")
    payload = {
        "tool": TOOL_NAME,
        "format_version": FORMAT_VERSION,
        "created_at": _now().isoformat(),
        "product": args.product,
        "round_number": round_number,
        "manual": {"path": str(manual.resolve()), "name": manual.stem, "format": manual.suffix.lower().lstrip(".")},
        "previous_manual": str(previous.resolve()) if previous else None,
        "references": [{"kind": "release_note", "path": str(p.resolve())} for p in release_notes]
        + [{"kind": "design_review", "path": str(p.resolve())} for p in design_reviews],
        "spec_documents": spec_names,
        "warnings": warnings,
        "judgment_values": [{"code": item.value, "label": JUDGMENT_LABELS_KO[item]} for item in ManualJudgment],
        "evidence_types": [item.value for item in EvidenceType],
        "summary": {
            "total_changes": len(changes),
            "functional_changes": len(functional),
            "image_changes": sum(1 for item in changes if item["image_change"]),
            "release_scope_total": len(release_findings),
            "release_scope_missing_suspected": missing,
            "prior_open_comments": len(open_prior),
        },
        "changes": changes,
        "release_findings": release_findings,
        "prior_comments": open_prior,
        "prior_comments_closed": closed_prior,
    }
    template = {
        "안내": (
            "변경마다 decision(판정 코드), confidence(0~1), evidence(근거 사양 조각 chunk_id 와 type)를 채운다. "
            "문제가 있으면 problem, recommended_manual_text, qa_comment 를 쓴다. 근거가 없으면 evidence 를 비워 둔다. "
            "prior_comments 의 status 는 QA 가 정한다."
        ),
        "judgment_values": {item.value: JUDGMENT_LABELS_KO[item] for item in ManualJudgment},
        "decisions": [
            {
                "change_id": item["change_id"],
                "kind": item["kind"],
                "text": item["text"][:200],
                "image_change": item["image_change"],
                "candidate_ids": [candidate["chunk_id"] for candidate in item["srs_candidates"]],
                "decision": "",
                "confidence": None,
                "evidence": [],
                "problem": "",
                "recommended_manual_text": "",
                "qa_comment": "",
            }
            for item in functional
        ],
        "prior_comments": [
            {
                "comment_id": item.get("comment_id"),
                "comment_text": item.get("comment_text", ""),
                "change_text": (item.get("change_text") or "")[:120],
                "suggestion": item["resolution_suggestion"]["label"],
                "status": item.get("status", CommentStatus.OPEN.value),
            }
            for item in open_prior
        ],
    }
    run_dir = _unique_run_dir(out_root, manual)
    _write_json(run_dir / "changes.json", payload)
    _write_json(run_dir / "decisions.json", template)

    print(f"결과 폴더: {run_dir}")
    print(f"변경 {len(changes)}건 · 분석 대상 {len(functional)}건 · 이미지 변경 {payload['summary']['image_changes']}건")
    if release_findings:
        print(f"Release 범위 {len(release_findings)}건 · 누락 의심 {missing}건")
    if open_prior:
        print(f"이전 회차 미해결 지적사항 {len(open_prior)}건")
    print(f"사양서 {len(spec_names)}개 · 조각 {len(chunks)}개")
    for warning in warnings:
        print(f"경고: {warning}")
    return EXIT_OK


# ---------------------------------------------------------------- finish


def _validate_decisions(changes: dict, decisions: dict) -> tuple[list[str], list[str], dict[str, dict]]:
    """반환값: (빈 판정 변경 번호, 그 밖의 오류, 변경 번호별 판정 줄)."""
    functional = {item["change_id"]: item for item in changes["changes"] if item["functional"]}
    all_ids = {item["change_id"] for item in changes["changes"]}
    allowed = {item.value for item in ManualJudgment}
    evidence_types = {item.value for item in EvidenceType}
    rows: dict[str, dict] = {}
    errors: list[str] = []
    for row in decisions.get("decisions") or []:
        change_id = str(row.get("change_id") or "")
        if change_id not in all_ids:
            errors.append(f"{change_id or '(빈 번호)'}: changes.json 에 없는 변경 번호입니다.")
            continue
        if change_id not in functional:
            errors.append(f"{change_id}: 단순 변경이라 판정하지 않습니다. 이 줄을 지우세요.")
            continue
        if change_id in rows:
            errors.append(f"{change_id}: 같은 변경의 판정이 두 번 있습니다.")
            continue
        rows[change_id] = row
    empty = [change_id for change_id in functional if not str((rows.get(change_id) or {}).get("decision") or "").strip()]
    for change_id, row in rows.items():
        decision = str(row.get("decision") or "").strip()
        if not decision:
            continue
        if decision not in allowed:
            errors.append(f"{change_id}: 판정 값 '{decision}' 은 목록에 없습니다. 가능한 값: {', '.join(sorted(allowed))}")
        confidence = row.get("confidence")
        if isinstance(confidence, bool) or not isinstance(confidence, (int, float)) or not 0 <= confidence <= 1:
            errors.append(f"{change_id}: 확신도(confidence)는 0~1 사이 숫자여야 합니다 (지금: {confidence!r}).")
        candidates = {item["chunk_id"] for item in functional[change_id]["srs_candidates"]}
        for evidence in row.get("evidence") or []:
            chunk_id = evidence if isinstance(evidence, str) else (evidence or {}).get("chunk_id")
            evidence_type = "DIRECT_SPEC" if isinstance(evidence, str) else (evidence or {}).get("type") or "DIRECT_SPEC"
            if chunk_id not in candidates:
                errors.append(f"{change_id}: 근거 '{chunk_id}' 는 이 변경에 보낸 사양 조각 후보가 아닙니다.")
            if evidence_type not in evidence_types:
                errors.append(f"{change_id}: 근거 종류 '{evidence_type}' 는 목록에 없습니다.")
    prior_ids = {item.get("comment_id") for item in changes.get("prior_comments") or []}
    statuses = {item.value for item in CommentStatus}
    for item in decisions.get("prior_comments") or []:
        if item.get("comment_id") not in prior_ids:
            errors.append(f"이전 지적사항 {item.get('comment_id')}: changes.json 에 없는 번호입니다.")
        elif item.get("status") not in statuses:
            errors.append(f"이전 지적사항 {item.get('comment_id')}: 상태 '{item.get('status')}' 는 목록에 없습니다.")
    return empty, errors, rows


def _judgment_for(change: dict, row: dict) -> ManualChangeJudgment:
    candidates = {item["chunk_id"]: item for item in change["srs_candidates"]}
    evidence: list[Evidence] = []
    for item in row.get("evidence") or []:
        chunk_id = item if isinstance(item, str) else item.get("chunk_id")
        evidence_type = "DIRECT_SPEC" if isinstance(item, str) else item.get("type") or "DIRECT_SPEC"
        candidate = candidates[chunk_id]
        evidence.append(Evidence(
            type=EvidenceType(evidence_type), srs_title=candidate["heading"], source_file=candidate["document"],
            page=int(candidate["page"] or 0), section=candidate["heading"], chunk_id=chunk_id,
        ))
    judgment = ManualChangeJudgment(
        decision=ManualJudgment(str(row["decision"]).strip()),
        confidence=float(row["confidence"]),
        reason_codes=[str(code) for code in row.get("reason_codes") or []],
        problem=str(row.get("problem") or ""),
        recommended_manual_text=str(row.get("recommended_manual_text") or ""),
        qa_comment=str(row.get("qa_comment") or ""),
        evidence=evidence,
        needs_human_review=bool(row.get("needs_human_review")),
    )
    apply_human_review_signals(judgment, change["kind"], change["review_required"], change["release_context"])
    if not change["srs_candidates"] and change["needs_human_review"] and not change["review_required"]:
        judgment.needs_human_review = True
        if SRS_NOT_AVAILABLE_REASON not in judgment.reason_codes:
            judgment.reason_codes.append(SRS_NOT_AVAILABLE_REASON)
    if change["needs_human_review"]:
        judgment.needs_human_review = True
    return judgment


def _location(change: dict) -> str:
    if change.get("source_page"):
        return f"PDF 페이지 {change['source_page']}"
    return f"{change.get('author') or '-'} · {change.get('date') or '-'} · 문단 #{change['paragraph_index']}"


def _write_excel(path: Path, changes: dict, judgments: dict[str, ManualChangeJudgment], prior: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "변경 판정"
    header = ["변경 번호", "변경 종류", "위치", "변경 글자", "분석 대상", "판정", "판정 코드", "확신도",
              "사람 확인", "판정 신호", "문제", "권장 문구", "QA Comment", "근거"]
    sheet.append(header)
    for change in changes["changes"]:
        judgment = judgments.get(change["change_id"])
        if judgment is None:
            sheet.append([change["change_id"], change["kind_label"], _location(change), change["text"], "아니요 (단순 변경)",
                          "-", "", None, "", "", "", "", "", ""])
            continue
        evidence = "\n".join(
            f"{item.type.value} · {item.source_file} p.{item.page} · {item.srs_title} ({item.chunk_id})" for item in judgment.evidence
        )
        sheet.append([
            change["change_id"], change["kind_label"], _location(change), change["text"], "예",
            JUDGMENT_LABELS_KO[judgment.decision], judgment.decision.value, round(judgment.confidence, 2),
            "필요" if judgment.needs_human_review else "", ", ".join(judgment.reason_codes),
            judgment.problem, judgment.recommended_manual_text, judgment.qa_comment, evidence,
        ])
    for column, width in zip("ABCDEFGHIJKLMN", (9, 14, 28, 60, 14, 16, 26, 8, 9, 28, 40, 40, 40, 50)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    findings = workbook.create_sheet("누락 의심")
    findings.append(["출처", "문서", "카테고리", "제목", "상태", "짝 변경", "검토 결과", "설명"])
    for item in changes.get("release_findings") or []:
        findings.append([item["source"], item["source_document"], item["category"], item["title"],
                         RELEASE_STATUS_LABELS.get(item["status"], item["status"]), item["matched_change_id"] or "",
                         item["result_status"], item["description"]])
    if prior:
        sheet_prior = workbook.create_sheet("이전 지적사항")
        sheet_prior.append(["지적 번호", "회차", "지적 문구", "이전 변경 글자", "참고 판정", "상태"])
        for item in prior:
            suggestion = (item.get("resolution_suggestion") or {}).get("label", "")
            sheet_prior.append([item.get("comment_id"), item.get("round_number"), item.get("comment_text"),
                                (item.get("change_text") or "")[:120], suggestion,
                                COMMENT_STATUS_LABELS.get(item.get("status"), item.get("status"))])
    workbook.save(str(path))


def run_finish(args: argparse.Namespace) -> int:
    run_dir = Path(args.run_dir)
    changes = _read_json(run_dir / "changes.json")
    decisions = _read_json(run_dir / "decisions.json")
    if changes.get("tool") != TOOL_NAME:
        raise InputError(f"이 도구가 만든 changes.json 이 아닙니다: {run_dir / 'changes.json'}")

    empty, errors, rows = _validate_decisions(changes, decisions)
    if empty or errors:
        if empty:
            print(f"빈 판정이 있습니다 ({len(empty)}건): {', '.join(empty)}")
        for error in errors:
            print(f"오류: {error}")
        print("decisions.json 을 고친 뒤 finish 를 다시 실행하세요. 결과 파일은 만들지 않았습니다.")
        return EXIT_DECISION_ERROR

    by_id = {item["change_id"]: item for item in changes["changes"]}
    judgments = {change_id: _judgment_for(by_id[change_id], row) for change_id, row in rows.items()}

    final_rows = []
    for change_id, judgment in judgments.items():
        change = by_id[change_id]
        reason = ""
        if judgment.ai_original_decision is not None:
            reason = judgment.problem
        final_rows.append({
            "change_id": change_id,
            "image_change": change["image_change"],
            "decision": judgment.decision.value,
            "decision_label": JUDGMENT_LABELS_KO[judgment.decision],
            "ai_original_decision": judgment.ai_original_decision.value if judgment.ai_original_decision else None,
            "confidence": judgment.confidence,
            "needs_human_review": judgment.needs_human_review,
            "reason_codes": judgment.reason_codes,
            "reason": reason,
        })
    decisions["final"] = final_rows
    _write_json(run_dir / "decisions.json", decisions)

    prior_status = {item.get("comment_id"): item.get("status") for item in decisions.get("prior_comments") or []}
    prior = []
    for item in changes.get("prior_comments") or []:
        updated = {key: value for key, value in item.items() if key != "resolution_suggestion"}
        new_status = prior_status.get(item.get("comment_id"), item.get("status"))
        if new_status != item.get("status"):
            updated["status"] = new_status
            updated["status_changed_in"] = run_dir.name
        updated["resolution_suggestion"] = item.get("resolution_suggestion")
        prior.append(updated)

    excel_path = run_dir / "manual_review.xlsx"
    _write_excel(excel_path, changes, judgments, prior)

    manual = Path(changes["manual"]["path"])
    product = changes["product"]
    round_number = int(changes.get("round_number", 0))
    comment_rows = [
        {
            "id": change_id,
            "functional": True,
            "decision": judgment.decision.value,
            "qa_decision": None,
            "qa_note": "",
            "ai_judgment": judgment.model_dump(mode="json"),
            "paragraph_index": by_id[change_id]["paragraph_index"],
            "change_index_in_paragraph": by_id[change_id]["change_index_in_paragraph"],
        }
        for change_id, judgment in judgments.items()
    ]
    inserted: list[dict] = []
    comment_docx = None
    if changes["manual"]["format"] == "docx":
        if not manual.is_file():
            raise InputError(f"원본 매뉴얼을 찾을 수 없습니다: {manual}")
        comment_docx = run_dir / f"{manual.stem}_QA_Comment.docx"
        inserted = insert_comments_detailed(manual, comment_rows, comment_docx, author=f"{product} QA AI")
        if not inserted:
            comment_docx.unlink(missing_ok=True)
            comment_docx = None

    new_comments = [
        {
            "comment_id": f"R{round_number}-{item['change']['id']}",
            "change_id": item["change"]["id"],
            "round_number": round_number,
            "comment_text": item["text"],
            "change_text": by_id[item["change"]["id"]]["text"],
            "kind": by_id[item["change"]["id"]]["kind"],
            "paragraph_index": by_id[item["change"]["id"]]["paragraph_index"],
            "decision": item["change"]["decision"],
            "status": CommentStatus.OPEN.value,
            "run": run_dir.name,
        }
        for item in inserted
    ]
    history = [{key: value for key, value in item.items() if key != "resolution_suggestion"} for item in prior]
    comments_payload = {
        "tool": TOOL_NAME,
        "format_version": FORMAT_VERSION,
        "created_at": _now().isoformat(),
        "product": product,
        "manual": changes["manual"]["name"],
        "round_number": round_number,
        "comments": list(changes.get("prior_comments_closed") or []) + history + new_comments,
    }
    _write_json(run_dir / "comments.json", comments_payload)

    counts: dict[str, int] = {}
    for judgment in judgments.values():
        label = JUDGMENT_LABELS_KO[judgment.decision]
        counts[label] = counts.get(label, 0) + 1
    print(f"판정 {len(judgments)}건: " + ", ".join(f"{label} {count}" for label, count in counts.items()))
    for row in final_rows:
        if row["ai_original_decision"]:
            print(f"{row['change_id']}: 이미지 변경이라 문제없음을 판정 불가로 바꿨습니다.")
    print(f"결과 Excel: {excel_path}")
    if comment_docx:
        print(f"Word Comment 파일: {comment_docx} (Comment {len(inserted)}건)")
    elif changes["manual"]["format"] == "docx":
        print("Word Comment 파일: 문제 항목이 없어 만들지 않았습니다.")
    else:
        print("Word Comment 파일: PDF 매뉴얼이라 만들지 않습니다.")
    print(f"이번 회차 지적사항: {run_dir / 'comments.json'} ({len(new_comments)}건)")
    return EXIT_OK


# ---------------------------------------------------------------- 명령줄


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="manual_review_local.py", description="Claude 대화로 돌리는 매뉴얼 개정 검증 (REQ-MANUAL-018)")
    commands = parser.add_subparsers(dest="command", required=True)

    extract = commands.add_parser("extract", help="변경을 뽑아 changes.json 과 decisions.json 틀을 만든다")
    extract.add_argument("--manual", required=True, help="개정 매뉴얼 (.docx 변경 추적 또는 .pdf)")
    extract.add_argument("--previous", help="이전 판 매뉴얼 .pdf (개정 매뉴얼이 PDF 일 때 필요)")
    extract.add_argument("--release-note", action="append", default=[], help="Release Note (.docx/.pdf). 여러 번 줄 수 있다")
    extract.add_argument("--design-review", action="append", default=[], help="설계검토보고서 (.docx/.pdf). 여러 번 줄 수 있다")
    extract.add_argument("--product", default="VXvue", help="제품 설정 이름 (기본 VXvue)")
    extract.add_argument("--prior-comments", help="앞 회차에 이 도구가 만든 comments.json")
    extract.add_argument("--max-candidates", type=int, default=DEFAULT_MAX_CANDIDATES, help="변경 한 건의 사양 조각 후보 수 (기본 6)")
    extract.add_argument("--out-root", help="결과 폴더를 만들 곳 (기본 output/manual_review_local)")
    extract.add_argument("--knowledge-root", help="지식 사본(data/product_knowledge)이 있는 프로젝트 폴더 (기본 이 저장소)")

    finish = commands.add_parser("finish", help="decisions.json 을 검사하고 결과 파일을 만든다")
    finish.add_argument("--run-dir", required=True, help="extract 가 만든 결과 폴더")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        if args.command == "extract":
            return run_extract(args)
        return run_finish(args)
    except InputError as exc:
        print(f"입력 오류: {exc}")
        return EXIT_INPUT_ERROR
