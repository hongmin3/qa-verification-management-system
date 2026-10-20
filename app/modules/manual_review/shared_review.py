"""Reusable review work files; no AI calls, product DB writes or source edits."""
from __future__ import annotations

import bisect
import hashlib
import html
import json
import re
import shutil
import tempfile
import threading
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

from .source_inventory import MAX_FILE_BYTES, extract_inventory
from .prose_quality import validate_report_prose
from .report_explanations import EXPLANATION_FIELDS, LANGUAGE_CONTEXTS, evidence_text, validate_explanations

_LOCK = threading.RLock()
DECISIONS = ('수정 필수', '추가 필요', '보강 권장', '사양 확인 필요', '유지', '타 문서 위임 적절')
FINDING_KEYS = {'id', 'title', 'decision', 'priority', 'paragraphs', 'quote_paragraphs', 'proposed', 'basis',
                'coverage', 'reason', 'location', 'kind', 'note'} | EXPLANATION_FIELDS


def _write(path, value):
    path = Path(path)
    with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=path.parent, delete=False) as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        temporary = Path(stream.name)
    temporary.replace(path)


def _read(path):
    return json.loads(Path(path).read_text('utf-8'))


def _identifier(value):
    if not isinstance(value, str) or not re.fullmatch('[0-9a-f]{64}', value):
        raise ValueError('자료 ID 형식을 확인하세요.')
    return value


def _digest(data):
    return hashlib.sha256(data).hexdigest()


def _normalize(text):
    return ''.join(c for c in unicodedata.normalize('NFKC', text).casefold() if unicodedata.category(c)[0] in 'LNM')


def map_pdf(inventory, pdf_bytes, basis):
    """Normalized text is used only for matching, never as quoted source text."""
    import fitz
    with fitz.open(stream=pdf_bytes, filetype='pdf') as pdf:
        if pdf.needs_pass or not 0 < len(pdf) <= 2000:
            raise ValueError('대응 PDF의 암호와 페이지 수를 확인하세요.')
        page_texts = [page.get_text() for page in pdf]
        pages = [_normalize(text) for text in page_texts]
    footer_numbers = {}
    for index, text in enumerate(page_texts, 1):
        footer = re.search(r'Page\s+(\d+)\s+of\s+(\d+)', text)
        if footer:
            footer_numbers[str(index)] = int(footer[1])
    starts = [0]
    for page in pages:
        starts.append(starts[-1] + len(page))
    joined = ''.join(pages)
    candidates, anchors, result = {}, {}, {}
    for para in inventory['paragraphs']:
        number = para['number']
        if basis == 'final' and not para['text'] and para['deleted']:
            result[str(number)] = {'status': 'deleted_in_final', 'pages': []}
            continue
        fields = ('text',) if basis == 'final' else ('raw', 'text', 'before')
        for field in fields:
            query = _normalize(para[field])
            if len(query) < 6:
                continue
            positions, start = [], 0
            while len(positions) < 500:
                position = joined.find(query, start)
                if position < 0:
                    break
                positions.append((position, position + len(query)))
                start = position + len(query)
            if positions:
                candidates[number] = positions
                if len(positions) == 1:
                    anchors[number] = positions[0]
                break
    # Reject unique matches that contradict source order rather than picking an
    # apparently unique occurrence in a header, footer or another section.
    keys = sorted(anchors)
    prefix, suffix = {}, {}
    highest, lowest = -1, len(joined) + 1
    for number in keys:
        prefix[number] = highest
        highest = max(highest, anchors[number][0])
    for number in reversed(keys):
        suffix[number] = lowest
        lowest = min(lowest, anchors[number][0])
    blocked = {n for n in keys if anchors[n][0] < prefix[n] or anchors[n][0] > suffix[n]}
    anchors = {n: value for n, value in anchors.items() if n not in blocked}
    for _ in range(3):
        keys, added = sorted(anchors), 0
        for number, positions in candidates.items():
            if number in anchors or number in blocked:
                continue
            at = bisect.bisect_left(keys, number)
            low = anchors[keys[at - 1]][0] if at else 0
            high = anchors[keys[at]][1] if at < len(keys) else len(joined)
            bounded = [x for x in positions if low <= x[0] < high]
            if len(bounded) == 1:
                anchors[number] = bounded[0]
                added += 1
        if not added:
            break
    keys = sorted(anchors)
    for para in inventory['paragraphs']:
        number = para['number']
        if str(number) in result:
            continue
        positions = [anchors[number]] if number in anchors else candidates.get(number, [])
        found = sorted({page for pos in positions for page in range(
            bisect.bisect_right(starts, pos[0]), bisect.bisect_right(starts, pos[1] - 1) + 1)})
        if number in anchors:
            result[str(number)] = {'status': 'verified_text', 'pages': found}
            continue
        at = bisect.bisect_left(keys, number)
        # Short labels are accepted only between two known, correctly ordered anchors.
        if number not in blocked and 0 < at < len(keys):
            low, high = anchors[keys[at - 1]][0], anchors[keys[at]][1]
            query = _normalize(para['text'])
            hits, cursor = [], low
            if len(query) >= 2 and low < high:
                while cursor < high:
                    pos = joined.find(query, cursor, high)
                    if pos < 0:
                        break
                    hits.append(bisect.bisect_right(starts, pos))
                    cursor = pos + len(query)
            if len(set(hits)) == 1:
                result[str(number)] = {'status': 'verified_context', 'pages': sorted(set(hits))}
                continue
        result[str(number)] = {'status': 'ambiguous' if positions else 'unverified', 'pages': []}
    return {'basis': basis, 'sha256': _digest(pdf_bytes), 'physical_pages': len(pages), 'paragraphs': result,
            'page_number_basis': 'PDF physical page', 'printed_footer_numbers': footer_numbers,
            'footer_disagreements': [int(p) for p, label in footer_numbers.items() if int(p) != label]}


class ReviewStore:
    def __init__(self, root):
        self.base = Path(root).resolve() / 'data/manual_review_tools'

    def _folder(self, source_id):
        return self.base / 'sources' / _identifier(source_id)

    def source_path(self, source_id):
        folder = self._folder(source_id)
        try:
            meta = _read(folder / 'source.json')
        except (OSError, ValueError) as exc:
            raise ValueError('검토 자료를 찾을 수 없습니다. 업로드 또는 등록 문서 선택을 확인하세요.') from exc
        path = folder / ('source.' + meta['kind'])
        if not path.is_file() or _digest(path.read_bytes()) != meta['sha256']:
            raise ValueError('검토 자료의 원본 해시가 다릅니다.')
        return path

    def add(self, data, filename):
        name = str(filename).replace('\\', '/').rsplit('/', 1)[-1]
        suffix = Path(name).suffix.lower()
        if suffix not in ('.docx', '.pdf') or not data or len(data) > MAX_FILE_BYTES:
            raise ValueError('128MB 이하의 DOCX 또는 PDF 파일이 필요합니다.')
        source_id = _digest(data)
        folder = self._folder(source_id)
        with _LOCK:
            if not (folder / 'source.json').exists():
                staging = self.base / 'staging'
                staging.mkdir(parents=True, exist_ok=True)
                with tempfile.TemporaryDirectory(dir=staging) as directory:
                    temporary = Path(directory) / ('source' + suffix)
                    temporary.write_bytes(data)
                    inventory = extract_inventory(temporary)
                folder.mkdir(parents=True, exist_ok=True)
                target = folder / ('source' + suffix)
                target.write_bytes(data)
                meta = {'source_id': source_id, 'sha256': source_id, 'kind': suffix[1:], 'name': name,
                        'created_at': datetime.now(timezone.utc).isoformat(), 'pdfs': {}}
                _write(folder / 'inventory.json', inventory)
                _write(folder / 'source.json', meta)
            self.source_path(source_id)
            return _read(folder / 'source.json')

    def inventory(self, source_id):
        with _LOCK:
            self.source_path(source_id)
            folder = self._folder(source_id)
            meta, inventory = _read(folder / 'source.json'), _read(folder / 'inventory.json')
            inventory.update(source_id=source_id, name=meta['name'], pdfs=meta['pdfs'])
            page_maps = {basis: _read(folder / record['map'])['paragraphs'] for basis, record in meta['pdfs'].items()}
            for para in inventory['paragraphs']:
                para['pages'] = {basis: mapping.get(str(para['number']), {}) for basis, mapping in page_maps.items()}
            return inventory

    def attach_pdf(self, source_id, data, basis):
        if basis not in ('saved', 'final') or not data or len(data) > MAX_FILE_BYTES:
            raise ValueError('대응 PDF 종류와 크기를 확인하세요.')
        with _LOCK:
            inventory = self.inventory(source_id)
            mapping = map_pdf(inventory, data, basis)
            previous = self._folder(source_id)
            meta = _read(previous / 'source.json')
            filename = basis + '-' + mapping['sha256']
            meta['pdfs'][basis] = {'file': filename + '.pdf', 'map': filename + '.json',
                                  'sha256': mapping['sha256'], 'physical_pages': mapping['physical_pages'],
                                  'page_number_basis': mapping['page_number_basis'], 'footer_disagreements': mapping['footer_disagreements']}
            # A source ID fixes both original bytes and PDF display settings.
            # Later uploads cannot silently change another user's review pages.
            new_id = _digest(json.dumps({'sha256': meta['sha256'], 'pdfs': meta['pdfs']}, sort_keys=True).encode())
            folder = self._folder(new_id)
            if not (folder / 'source.json').exists():
                folder.mkdir(parents=True, exist_ok=True)
                for path in previous.iterdir():
                    if path.name != 'source.json' and path.is_file():
                        shutil.copyfile(path, folder / path.name)
                (folder / (filename + '.pdf')).write_bytes(data)
                _write(folder / (filename + '.json'), mapping)
                meta['source_id'] = new_id
                _write(folder / 'source.json', meta)
            return dict(mapping, source_id=new_id)

    def pdf_path(self, source_id, basis):
        path = self.source_path(source_id)
        if basis == 'source' and path.suffix == '.pdf':
            return path
        meta = _read(path.parent / 'source.json')
        record = meta['pdfs'].get(basis)
        if not record:
            raise ValueError('해당 표시 기준의 PDF가 없습니다.')
        return path.parent / record['file']


def read_review(store, source_id, view='summary', offset=0, limit=20, paragraphs=None):
    if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 40:
        raise ValueError('조회 시작 위치와 개수(1~40)를 확인하세요.')
    data = store.inventory(source_id)
    if view == 'summary':
        return {key: value for key, value in data.items() if key not in ('paragraphs', 'comments', 'comment_threads')} | {
            'paragraph_count': len(data['paragraphs']), 'comment_count': len(data['comments']),
            'changed_paragraphs': sum(bool(p['inserted'] or p['deleted']) for p in data['paragraphs']),
            'instruction': 'paragraphs·comments를 끝까지 조회하고, 현재 문구와 삭제 이력을 구분하세요. 현재 문구는 바꾸어 쓰지 마세요.'}
    if view not in ('paragraphs', 'comments'):
        raise ValueError('조회 종류를 확인하세요.')
    rows = data[view]
    if paragraphs is not None:
        if view != 'paragraphs' or not isinstance(paragraphs, list) or not 1 <= len(paragraphs) <= 40 or any(type(n) is not int or not 1 <= n <= len(rows) for n in paragraphs):
            raise ValueError('원본 문단 번호를 1~40개 지정하세요.')
        selected, end = [rows[n - 1] for n in paragraphs], len(paragraphs)
        total, truncated = end, False
    else:
        selected, end = rows[offset:offset + limit], min(offset + limit, len(rows))
        total, truncated = len(rows), end < len(rows)
    from app.core.security_filter import mask_text
    for row in selected:
        if 'text' in row:
            row['quote_transport_verbatim'] = mask_text(row['text'])[0] == row['text']
    return {'source_id': source_id, 'source_sha256': data['sha256'], 'view': view, 'results': selected,
            'total': total, 'truncated': truncated, 'next_offset': end if truncated else None,
            'quote_note': '조회는 민감정보 보호를 적용합니다. quote_transport_verbatim이 false이면 검색용 원문과 다릅니다. 생성 보고서의 현재 문구는 서버 원본에서 직접 채웁니다.'}


def _page_text(paragraphs, basis, source_kind):
    found, missing, deleted = set(), False, False
    for para in paragraphs:
        if source_kind == 'pdf' and basis == 'saved':
            found.add(para['physical_page'])
            continue
        record = para.get('pages', {}).get(basis, {})
        found.update(record.get('pages', []))
        if record.get('status') == 'deleted_in_final':
            deleted = True
        elif not record.get('pages') and (para['text'] or para['images']):
            missing = True
    parts = [', '.join(map(str, sorted(found))) + '쪽'] if found else []
    if missing:
        parts.append('일부 쪽수 미확인' if found else '쪽수 미확인')
    if deleted:
        parts.append('삭제 본문: 최종본 없음')
    return ' / '.join(parts) or '해당 없음'


def build_report(store, source_id, findings, title='매뉴얼 검토보고서', memo_review=None, release_review=None):
    if not isinstance(findings, list) or not 1 <= len(findings) <= 200:
        raise ValueError('검토 항목은 1~200개가 필요합니다.')
    if not isinstance(title, str) or not 1 <= len(title) <= 300:
        raise ValueError('보고서 제목을 확인하세요.')
    inventory = store.inventory(source_id)
    for appendix in (memo_review, release_review):
        if appendix is not None and (not isinstance(appendix, list) or len(appendix) > 1000 or any(not isinstance(row, dict) for row in appendix)):
            raise ValueError('메모·릴리즈 대조는 1000개 이하의 항목 목록으로 전달하세요.')
        if appendix and any(len(row) > 20 or any(not isinstance(k, str) or len(k) > 100 or not isinstance(v, (str, int, float, bool)) or len(str(v)) > 12000 for k, v in row.items()) for row in appendix):
            raise ValueError('메모·릴리즈 항목은 20개 이하의 짧은 열과 텍스트 값으로 전달하세요.')
    by_number = {p['number']: p for p in inventory['paragraphs']}
    rows, ids = [], set()
    for finding in findings:
        if not isinstance(finding, dict) or set(finding) - FINDING_KEYS:
            raise ValueError('검토 항목의 허용되지 않은 필드입니다. 현재 문구는 원본에서 채웁니다.')
        for field in ('id', 'title', 'decision', 'priority', 'proposed', 'basis', 'coverage', 'reason'):
            if not isinstance(finding.get(field), str) or not finding[field].strip() or len(finding[field]) > 12000:
                raise ValueError(f'검토 항목의 {field} 값을 확인하세요.')
        if finding['decision'] not in DECISIONS or finding['priority'] not in ('높음', '보통', '낮음'):
            raise ValueError('판정 또는 우선순위 값을 확인하세요.')
        if finding['id'] in ids or len(finding['id']) > 80:
            raise ValueError('검토 항목 ID가 중복되거나 너무 깁니다.')
        ids.add(finding['id'])
        validate_explanations(finding)
        targets = finding.get('paragraphs')
        quote_numbers = finding.get('quote_paragraphs', targets)
        for numbers in (targets, quote_numbers):
            if not isinstance(numbers, list) or len(numbers) > 40 or any(type(n) is not int or n not in by_number for n in numbers):
                raise ValueError('원본에 있는 문단 번호를 40개 이하로 지정하세요.')
        kind = finding.get('kind', 'manual')
        if kind not in ('manual', 'proposal', 'source') or not targets and kind == 'manual':
            raise ValueError('원문 문단 또는 추가 제안·근거문서 구분이 필요합니다.')
        if kind == 'source':
            # 근거 문서의 문단 번호는 이 매뉴얼의 같은 번호 문단으로 해석되므로 받지 않는다.
            if targets or finding.get('quote_paragraphs'):
                raise ValueError(f'{finding["id"]}: 근거 문서 자체 검토 항목에는 검토 대상 매뉴얼의 문단 번호를 넣지 마세요. '
                                 '근거 문서의 이름과 위치는 location에 쓰고, 매뉴얼 위치는 별도 manual 항목으로 만드세요.')
            if not isinstance(finding.get('location'), str) or not finding['location'].strip():
                raise ValueError(f'{finding["id"]}.location: 근거 문서 자체 검토 항목은 근거 문서 이름과 위치가 필요합니다.')
        if any(by_number[n].get('unsupported_symbols') for n in quote_numbers):
            raise ValueError('인용 문단에 글꼴 특수 기호가 있습니다. 해당 기호를 원본에서 확인하고 인용 문단을 분리하세요.')
        quotes = [{'paragraph': n, 'text': by_number[n]['text'], 'has_strike': by_number[n]['has_strike']}
                  for n in dict.fromkeys(quote_numbers) if by_number[n]['text']]
        current = '\n\n'.join(q['text'] for q in quotes)
        if len(current) > 32767:
            raise ValueError('현재 문구가 Excel 셀 제한보다 깁니다. 인용 문단을 나누세요.')
        selected = [by_number[n] for n in targets]
        row = dict(finding, kind=kind, quotes=quotes, current_text=current,
                   document_order=min(targets) if targets else 100000000,
                   section=by_number[min(targets)]['section'] if targets else '근거 사양 자체 검토',
                   pages_saved=_page_text(selected, 'saved', inventory['source_kind']),
                   pages_final=_page_text(selected, 'final', inventory['source_kind']))
        rows.append(row)
    rows.sort(key=lambda r: (r['kind'] == 'source', r['document_order'], r['id']))
    validate_report_prose(title, rows, memo_review, release_review, inventory)
    key = _digest(json.dumps({'source_id': source_id, 'pdfs': inventory['pdfs'], 'title': title,
                              'findings': rows, 'memo_review': memo_review, 'release_review': release_review,
                              'format_version': 4}, ensure_ascii=False, sort_keys=True).encode())
    folder = store.base / 'reports' / key
    folder.mkdir(parents=True, exist_ok=True)
    payload = {'source_id': source_id, 'source_sha256': inventory['sha256'], 'source_name': inventory['name'], 'source_kind': inventory['source_kind'],
               'title': title, 'pdfs': inventory['pdfs'], 'findings': rows, 'warnings': inventory['warnings'],
               'memo_review': memo_review or [], 'release_review': release_review or []}
    with _LOCK:
        _write(folder / 'report.json', payload)
        _write_html(folder / 'manual_review.html', payload)
        _write_excel(folder / 'manual_review.xlsx', payload, inventory)
    return {'report_id': key, 'source_id': source_id, 'count': len(rows), 'directory': str(folder),
            'files': ['manual_review.html', 'manual_review.xlsx', 'report.json'], 'source_hash_unchanged': True}


def _write_html(path, payload):
    e = lambda text: html.escape(str(text))
    cards = []
    for row in payload['findings']:
        quotes = ''.join(f'<figure><figcaption>문단 {q["paragraph"]}'
                         + (' · 취소선 포함' if q['has_strike'] else '')
                         + f'</figcaption><pre>{e(q["text"])}</pre><button type="button" onclick="navigator.clipboard.writeText(this.previousElementSibling.textContent)">원문 복사</button></figure>' for q in row['quotes'])
        if not quotes:
            quotes = '<p>인용할 현재 문구 없음. 추가 제안·삭제 이력·근거문서 검토를 구분하여 확인하십시오.</p>'
        explanation = '<h3>현재 원문의 한국어 뜻</h3><p>' + e(row.get('current_meaning', '한국어 뜻 미작성')) + '</p>'
        if row.get('language_context'):
            explanation += '<p>언어 문맥: ' + e(LANGUAGE_CONTEXTS[row['language_context']]) + '</p>'
        evidence = ''.join('<figure><figcaption>' + e(item['source']) + ' · ' + e(item['location'])
                           + '</figcaption><pre class="evidence-quote">' + e(item['quote'])
                           + '</pre><p>한국어 뜻: ' + e(item['meaning']) + '</p><p>대조 관계 (검토자 해설): '
                           + e(item['relation']) + '</p></figure>' for item in row.get('evidence', []))
        explanation += '<h3>확인한 근거 (검토자 입력)</h3><p>' + e(row['basis']) + '</p>' + evidence
        details = ''.join(f'<dt>{label}</dt><dd>{e(row[field])}</dd>' for label, field in (
            ('대조 판단 (AI 분석)', 'reason'), ('수정안 또는 리뷰 요청', 'proposed'), ('공통 및 타 문서 대조', 'coverage')))
        if row.get('note'):
            details += '<dt>검토 설명</dt><dd>' + e(row['note']) + '</dd>'
        comment = ('<h3>한국어 검토 코멘트</h3><pre class="review-comment">' + e(row['short_comment'])
                   + '</pre><button type="button" data-copy="short-comment" onclick="copyReviewText(this)">한국어 검토 코멘트 복사</button>') if row.get('short_comment') else '<p>짧은 메모 미작성</p>'
        cards.append(f'<article id="{e(row["id"])}"><small>{e(row["id"])} · {e(row["decision"])} · {e(row["priority"])}</small>'
                     f'<h2>{e(row["title"])}</h2><p>{e(row.get("location") or row["section"])} · 문단 {e(row["paragraphs"])}</p>'
                     f'<p>원본/변경추적 표시: {e(row["pages_saved"])} · 최종본: {e(row["pages_final"])}</p>'
                     f'<h3>현재 문구 · 원문 그대로</h3>{quotes}{explanation}<dl>{details}</dl>{comment}</article>')
    appendix = ''
    for title, records in [('기존 메모 검토', payload['memo_review']), ('릴리즈 항목 대조', payload['release_review'])]:
        if records:
            appendix += '<h2>' + title + '</h2>'
            for record in records:
                appendix += '<article><dl>' + ''.join('<dt>' + e(k) + '</dt><dd>' + e(v) + '</dd>' for k, v in record.items()) + '</dl></article>'
    body = f'''<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{e(payload['title'])}</title>
<style>body{{font:15px/1.65 'Malgun Gothic',sans-serif;background:#f5f6f8;color:#182334}}main{{max-width:1100px;margin:auto;padding:24px}}article{{background:white;border:1px solid #d6dde7;padding:24px;margin:18px 0;break-inside:avoid}}h2{{font-size:20px}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit;margin:8px 0}}figure{{margin:14px 0;border-left:3px solid #607d9b;padding:12px}}figcaption,small{{color:#53647b}}dl{{display:grid;grid-template-columns:170px 1fr;gap:12px}}dd{{margin:0;white-space:pre-wrap;overflow-wrap:anywhere}}dt{{font-weight:bold}}button{{padding:5px 12px}}@media print{{button{{display:none}}}}</style><main><h1>{e(payload['title'])}</h1>
<p>{e(payload['source_name'])} · 검토 {len(payload['findings'])}건 · 매뉴얼 순서, 기존 ID 유지</p>
<p>현재 문구는 원본에서 직접 인용했습니다. 인용 안의 공백·기호를 보존하며, 문단 번호는 쪽수가 아닙니다. 쪽수는 대응 PDF의 물리 페이지입니다. 쪽수 미확인 항목은 절과 문단으로 찾으십시오.</p>
{''.join(cards)}{appendix}<p>원본 SHA-256: {payload['source_sha256']}</p></main>
<style>body{{overflow-wrap:anywhere}}dl{{grid-template-columns:170px minmax(0,1fr)}}@media(max-width:640px){{main{{padding:12px}}article{{padding:16px}}dl{{grid-template-columns:minmax(0,1fr);gap:6px}}dd{{margin-bottom:12px}}}}</style>
<script>async function copyReviewText(button){{const text=button.previousElementSibling.textContent;try{{if(navigator.clipboard && window.isSecureContext){{await navigator.clipboard.writeText(text);}}else{{const area=document.createElement('textarea');area.value=text;area.style.position='fixed';area.style.opacity='0';document.body.appendChild(area);area.select();const copied=document.execCommand('copy');area.remove();if(!copied)throw new Error('copy');}}button.textContent='복사 완료';}}catch(error){{button.textContent='복사 실패: 위 문구를 선택해 복사';}}}}</script></html>'''
    path.write_text(body, encoding='utf-8')


def _write_excel(path, payload, inventory):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    headers = ['ID', '판정', '우선순위', '절·문단 위치', '실제 쪽수 (원본/변경추적 표시)', '실제 쪽수 (최종본)',
               '검토 항목', '현재 문구 (원문 그대로)', '수정안 또는 리뷰 요청', '근거', '공통 및 타 문서 대조', '이유', '인용 문단', '검토 설명',
               '현재 원문의 한국어 뜻', '근거 원문과 해설', '한국어 검토 코멘트', '언어 문맥']
    findings = [[r['id'], r['decision'], r['priority'], (r.get('location') or r['section']) + ' · 문단 ' + str(r['paragraphs']),
                 r['pages_saved'], r['pages_final'], r['title'], r['current_text'], r['proposed'], r['basis'], r['coverage'], r['reason'],
                 ', '.join(str(q['paragraph']) for q in r['quotes']), r.get('note', ''), r.get('current_meaning', ''),
                 evidence_text(r.get('evidence', [])), r.get('short_comment', ''), LANGUAGE_CONTEXTS.get(r.get('language_context'), '')] for r in payload['findings']]
    comments = [[c['id'], str(c['paragraphs']), c['text'], '반영·해결 상태 별도 검토'] for c in inventory['comments']]
    changes = [[p['number'], p['section'], p['before'], p['text'], p['inserted'], p['deleted'], str(p['comment_ids'])]
               for p in inventory['paragraphs'] if p['inserted'] or p['deleted']]
    scope = [['대상', payload['source_name'], payload['source_id']],
             ['본문', str(len(inventory['paragraphs'])) + '개 문단', '추출 수집과 사양 전수검토 완료는 별개'],
             ['실제 쪽수', json.dumps(payload['pdfs'], ensure_ascii=False), '제공 PDF에서 대응한 위치만 표시'],
             ['원문 인용', '서버가 원본에서 직접 채움', '삭제 이력·해설은 현재 문구에 섞지 않음'],
             ['미수행', '제품 실행·사양 최종 승인', '\n'.join(payload['warnings'])]]
    tables = [('검토결과', headers, findings),
              ('메모검토', ['메모 ID', '연결 문단', '메모 원문', '상태'], comments),
              ('릴리즈대조', ['항목', '대조 결과'], [['별도 검토', '릴리즈 대조 판정은 검토자가 근거를 조사하여 입력합니다.']]),
              ('변경추적', ['문단', '절', '변경 거부 시 텍스트', '현재 표시 텍스트', '추가', '삭제', '메모 ID'], changes),
              ('조사범위', ['구분', '확인 범위', '설명 및 제한'], scope)]
    if payload['memo_review']:
        keys = list(dict.fromkeys(k for row in payload['memo_review'] for k in row))
        tables[1] = ('메모검토', keys, [[row.get(k, '') for k in keys] for row in payload['memo_review']])
    if payload['release_review']:
        keys = list(dict.fromkeys(k for row in payload['release_review'] for k in row))
        tables[2] = ('릴리즈대조', keys, [[row.get(k, '') for k in keys] for row in payload['release_review']])
    wb = Workbook()
    wb.remove(wb.active)
    for index, (name, columns, rows) in enumerate(tables):
        ws = wb.create_sheet(name)
        ws['A2'] = payload['title'] if index == 0 else name
        ws['A2'].font = Font(name='Malgun Gothic', size=15, bold=True)
        ws['A4'] = '원문 그대로 인용 · 매뉴얼 순서 · 문단 번호와 실제 쪽수 구분'
        for col, value in enumerate(columns, 1):
            cell = ws.cell(6, col, value)
            cell.fill = PatternFill('solid', fgColor='233A55')
            cell.font = Font(name='Malgun Gothic', bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws.column_dimensions[cell.column_letter].width = (12 if col <= 3 and index == 0 else 46)
        for row_number, values in enumerate(rows, 7):
            for col, value in enumerate(values, 1):
                value = str(value)
                if len(value) > 32767:
                    raise ValueError('Excel 셀 제한보다 긴 원문입니다. 검토 항목을 나누세요.')
                cell = ws.cell(row_number, col, value)
                cell.data_type = 's'  # Source strings beginning with '=' are literal text, never formulas.
                cell.font = Font(name='Malgun Gothic', size=10)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            lines = max((max(1, len(str(v)) // 35) + str(v).count('\n') for v in values), default=1)
            ws.row_dimensions[row_number].height = min(409, max(50, 15 * lines + 12))
        if rows:
            table = Table(displayName=f'Review_{index}', ref=f'A6:{ws.cell(6, len(columns)).column_letter}{len(rows) + 6}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
            ws.add_table(table)
        ws.freeze_panes = 'D7' if index == 0 else 'A7'
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[6].height = 34
    wb.save(path)
