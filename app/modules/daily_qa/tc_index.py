"""TC Excel 색인 (SPEC REQ-DAILY-003, REQ-DAILY-005).

`app/parsers/excel_parser.py` 는 TC ID 중심이라 SRS 번호 열을 읽지 않는다. 일일 점검은
TC 와 SRS 를 잇는 열(`Old SRS ID`, `SRS No`)이 핵심이므로 여기서 따로 읽는다.
원본 파일은 읽기 전용으로 연다.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from openpyxl import load_workbook

#: 정규화(영숫자 소문자) 이름 → 필드. 앞쪽이 우선한다.
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "srs": ("oldsrsid", "srsno", "srsid", "srs"),
    "tc_id": ("tcid", "oldid", "testcaseid", "id"),
    "title": ("title", "testitem"),
    "precondition": ("precondition",),
    "step": ("teststep", "stepdescription", "step", "steps"),
    "expected": ("expectedresult", "expectedreuslt", "expected"),
    "change": ("변경사항",),
}

#: SRS 번호 모양. Polarion ID(VP-123)와 Legacy 번호(03-10-05, 05-140-10)를 모두 잡는다.
SRS_TOKEN_RE = re.compile(r"\b(?:[A-Z]{2,}-\d+|\d{2}-\d{1,3}(?:-\d{1,3})+)\b")

HEADER_SCAN_ROWS = 12


@dataclass
class TcRow:
    workbook: str
    sheet: str
    row: int
    tc_id: str
    srs_refs: list[str] = field(default_factory=list)
    title: str = ""
    precondition: str = ""
    step: str = ""
    expected: str = ""

    def location(self) -> str:
        return f"{self.workbook} / {self.sheet} / {self.row}행"

    def as_dict(self) -> dict:
        data = asdict(self)
        data["location"] = self.location()
        return data


def _normalize(value: object) -> str:
    return re.sub(r"[^0-9a-z가-힣]", "", str(value or "").casefold())


def _detect(header: tuple) -> dict[str, int] | None:
    normalized = [_normalize(cell) for cell in header]
    columns: dict[str, int] = {}
    for name, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                columns[name] = normalized.index(alias)
                break
    # SRS 열과 Title/Step 중 하나가 함께 있어야 TC 표로 본다.
    if "srs" in columns and ("title" in columns or "step" in columns):
        return columns
    return None


def srs_tokens(value: str) -> list[str]:
    return list(dict.fromkeys(SRS_TOKEN_RE.findall(value or "")))


def read_workbook(path: Path) -> list[TcRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[TcRow] = []
    try:
        for sheet in workbook.worksheets:
            header_row = None
            columns = None
            for index, values in enumerate(sheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), start=1):
                columns = _detect(values)
                if columns:
                    header_row = index
                    break
            if not columns or header_row is None:
                continue

            def cell(values: tuple, name: str) -> str:
                position = columns.get(name)
                if position is None or position >= len(values):
                    return ""
                return str(values[position] or "").strip()

            for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
                title, step = cell(values, "title"), cell(values, "step")
                if not (title or step):
                    continue
                rows.append(
                    TcRow(
                        workbook=path.name,
                        sheet=sheet.title,
                        row=row_number,
                        tc_id=cell(values, "tc_id"),
                        srs_refs=srs_tokens(cell(values, "srs")),
                        title=title,
                        precondition=cell(values, "precondition"),
                        step=step,
                        expected=cell(values, "expected"),
                    )
                )
    finally:
        workbook.close()
    return rows


def build_index(paths: list[Path]) -> tuple[list[TcRow], list[str]]:
    """여러 TC 파일을 읽는다. 읽지 못한 파일은 이유와 함께 돌려준다 (숨기지 않는다)."""
    rows: list[TcRow] = []
    errors: list[str] = []
    for path in paths:
        try:
            rows.extend(read_workbook(path))
        except Exception as exc:  # 한 파일이 깨져도 나머지는 색인한다
            errors.append(f"{path.name}: {type(exc).__name__}")
    return rows, errors


def by_srs(rows: list[TcRow]) -> dict[str, list[TcRow]]:
    index: dict[str, list[TcRow]] = {}
    for row in rows:
        for ref in row.srs_refs:
            index.setdefault(ref, []).append(row)
    return index
