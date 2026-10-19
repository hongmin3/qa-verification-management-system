"""Product-scoped verification sheet selection. Never silently use another version."""
from __future__ import annotations

import hashlib
from openpyxl import load_workbook
from app.modules.daily_qa.tc_index import read_workbook


def versions(paths):
    options, errors = [], []
    for path in paths:
        try:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for name in book.sheetnames:
                    key = hashlib.sha256(f'{path.name}\0{name}'.encode()).hexdigest()[:24]
                    options.append({'key': key, 'workbook': path.name, 'sheet': name})
            finally:
                book.close()
        except Exception as exc:
            errors.append(f'{path.name}: {type(exc).__name__}')
    return options, errors


def selected(store, product, options):
    key = store.get_state(f'{product}:verification_sheet')
    if key:
        matches = [item for item in options if item['key'] == key]
        return matches[0] if len(matches) == 1 else None
    defaults = [item for item in options if item['sheet'] == 'V1.1.0']
    return defaults[0] if len(defaults) == 1 else None


def save(store, product, key, options):
    matches = [item for item in options if item['key'] == key]
    if len(matches) != 1:
        raise ValueError('등록된 Checklist 파일과 시트를 선택해 주세요.')
    store.set_state(f'{product}:verification_sheet', key)
    return matches[0]


def selected_rows(paths, selection):
    if not selection:
        return [], ['검증 버전의 Checklist 시트를 선택해 주세요.']
    matches = [path for path in paths if path.name == selection['workbook']]
    if len(matches) != 1:
        return [], ['선택한 Checklist 파일을 하나로 확인할 수 없습니다.']
    try:
        rows = [row for row in read_workbook(matches[0]) if row.sheet == selection['sheet']]
    except Exception as exc:
        return [], [f"{selection['workbook']}: {type(exc).__name__}"]
    return (rows, []) if rows else ([], ['선택한 시트에서 읽을 수 있는 TC가 없습니다.'])


def for_choice(choice):
    from app.modules.daily_qa.pipeline import default_inputs
    marker = (choice.cfg.product_profile.checklist or {}).get('template_name_contains') or ''
    inputs = default_inputs(choice.cfg.product, choice.cfg.root, marker)
    paths = checklist_paths(inputs, choice.cfg.product_profile)
    options, errors = versions(paths)
    return {'options': options, 'selected': selected(choice.store, choice.cfg.slug, options),
            'errors': errors, 'enabled': choice.cfg.savings_enabled,
            'five_hour_percent': choice.cfg.usage_guard_five_hour_percent, 'seven_day_percent': choice.cfg.usage_guard_seven_day_percent}


def checklist_paths(inputs, profile):
    marker = (profile.checklist or {}).get('template_name_contains') or ''
    paths = [path for path in inputs.tc_paths if marker and marker in path.name]
    return paths or ([inputs.checklist_template] if inputs.checklist_template else [])
